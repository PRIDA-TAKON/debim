"""
CLI interface for debim using standard library argparse and pure-Python formatting
"""

import argparse
import re
import sys
from pathlib import Path
from typing import Optional, List, Dict, Any
import yaml
from pydantic import ValidationError

if sys.platform == "win32":
    try:
        if hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(encoding="utf-8")
        if hasattr(sys.stderr, "reconfigure"):
            sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

from debim.version import __version__, check_and_notify_updates
from debim.compiler import compile_to_ifc
from debim.ingest import ingest_dxf_file
from debim.bsdd import validate_manifest_psets, PropertyValidationStatus
from debim.scaffold import scaffold_element
from debim.cost import estimate_cost, generate_cost_template, load_price_catalog
from debim.qto import calculate_qto
from debim.resolver import resolve_manifest
from debim.schema import ProjectManifest, load_manifest, export_json_schema
from debim.loader import resolve_default_manifest_path
from debim.compliance import check_topological_continuity, detect_mep_structural_clashes
from debim.viewer import generate_viewer_html, serve_viewer
from debim.modular import bundle_manifest, split_manifest
from debim.spec.registry import SpecRegistryClient
from debim.docs.exemplar import (
    generate_exemplar_element,
    generate_exemplar_manifest,
    get_supported_exemplar_classes,
)
from debim.docs.thai_specs import get_thai_spec_clause, THAI_SPEC_LIBRARY

# --- Pure-Python Rich-like formatting helpers ---

_STYLE_MAP = {
    "bold": "\033[1m",
    "dim": "\033[2m",
    "red": "\033[31m",
    "green": "\033[32m",
    "yellow": "\033[33m",
    "blue": "\033[34m",
    "magenta": "\033[35m",
    "cyan": "\033[36m",
    "white": "\033[37m",
    "reset": "\033[0m",
}

def _color_tag_replacer(match: re.Match) -> str:
    tag = match.group(1).strip()
    if tag.startswith("/"):
        return _STYLE_MAP["reset"]

    parts = tag.split()
    codes = []
    for part in parts:
        if part in _STYLE_MAP:
            codes.append(_STYLE_MAP[part])
        elif part.startswith("#"):
            # Hex color or approximation, fallback to cyan/bold if hex
            codes.append("\033[36m")
    return "".join(codes) if codes else ""

def format_markup(text: str) -> str:
    """Convert Rich markup like [bold green]text[/bold green] to ANSI escape codes or clean text."""
    if not sys.stdout.isatty():
        # Strip markup if not tty
        return re.sub(r"\[/?[^\]]+\]", "", text)
    formatted = re.sub(r"\[([a-zA-Z0-9_#/ -]+)\]", _color_tag_replacer, text)
    return formatted

def strip_markup(text: str) -> str:
    """Remove markup tags and ANSI codes for calculating string lengths."""
    clean = re.sub(r"\[/?[^\]]+\]", "", text)
    clean = re.sub(r"\033\[[0-9;]*[a-zA-Z]", "", clean)
    return clean

class PureConsole:
    def print(self, *args, end: str = "\n", file=None):
        out_file = file or sys.stdout
        msg = " ".join(str(a) for a in args)
        out_file.write(format_markup(msg) + end)
        out_file.flush()

console = PureConsole()

class PurePanel:
    def __init__(self, renderable: str, title: Optional[str] = None, style: Optional[str] = None):
        self.renderable = renderable
        self.title = title
        self.style = style

    def __str__(self):
        lines = str(self.renderable).splitlines()
        max_len = max((len(strip_markup(l)) for l in lines), default=0)
        if self.title:
            max_len = max(max_len, len(strip_markup(self.title)))
        box_width = max_len + 4

        border_style = f"[{self.style}]" if self.style else ""
        border_end = f"[/{self.style}]" if self.style else ""

        top_title = f" {self.title} " if self.title else ""
        top_bar = "─" * (box_width - 2 - len(strip_markup(top_title)))
        res = [f"{border_style}┌{top_title}{top_bar}┐{border_end}"]

        for line in lines:
            padding = box_width - 4 - len(strip_markup(line))
            res.append(f"{border_style}│{border_end} {line}{' ' * padding} {border_style}│{border_end}")

        res.append(f"{border_style}└{'─' * (box_width - 2)}┘{border_end}")
        return "\n".join(res)

Panel = PurePanel

class PureTable:
    def __init__(
        self,
        title: Optional[str] = None,
        show_header: bool = True,
        header_style: Optional[str] = None,
        show_lines: bool = False,
    ):
        self.title = title
        self.show_header = show_header
        self.header_style = header_style or "bold cyan"
        self.columns: List[Dict[str, Any]] = []
        self.rows: List[List[str]] = []

    def add_column(self, name: str, style: Optional[str] = None, justify: str = "left", width: Optional[int] = None):
        self.columns.append({"name": name, "style": style or "", "justify": justify, "width": width})

    def add_row(self, *args):
        self.rows.append([str(a) for a in args])

    def __str__(self):
        col_widths = []
        for i, col in enumerate(self.columns):
            w = len(strip_markup(col["name"])) if self.show_header else 0
            for row in self.rows:
                if i < len(row):
                    w = max(w, len(strip_markup(row[i])))
            col_widths.append(w)

        lines = []
        if self.title:
            lines.append(f"[{self.header_style}]{self.title}[/{self.header_style}]")

        if self.show_header:
            header_cells = []
            for i, col in enumerate(self.columns):
                w = col_widths[i]
                c_text = col["name"]
                pad = w - len(strip_markup(c_text))
                if col["justify"] == "right":
                    cell = " " * pad + c_text
                elif col["justify"] == "center":
                    left = pad // 2
                    right = pad - left
                    cell = " " * left + c_text + " " * right
                else:
                    cell = c_text + " " * pad
                header_cells.append(cell)
            header_str = "  ".join(header_cells)
            lines.append(f"[{self.header_style}]{header_str}[/{self.header_style}]")
            lines.append("─" * len(strip_markup(header_str)))

        for row in self.rows:
            row_cells = []
            for i, col in enumerate(self.columns):
                w = col_widths[i]
                val = row[i] if i < len(row) else ""
                pad = w - len(strip_markup(val))
                if col["justify"] == "right":
                    cell = " " * pad + val
                elif col["justify"] == "center":
                    left = pad // 2
                    right = pad - left
                    cell = " " * left + val + " " * right
                else:
                    cell = val + " " * pad

                if col["style"] and not re.search(r"\[/?[a-zA-Z0-9_# -]+\]", cell):
                    cell = f"[{col['style']}]{cell}[/{col['style']}]"
                row_cells.append(cell)
            lines.append("  ".join(row_cells))

        return "\n".join(lines)

Table = PureTable


DEBIM_BANNER = (
    "\n"
    "[bold #00ffff]  ███╗   [/bold #00ffff]         [bold #00ffff]██████╗ [/bold #00ffff][bold #22d3ee]███████╗[/bold #22d3ee][bold #38bdf8]██████╗ [/bold #38bdf8][bold #60a5fa]██╗[/bold #60a5fa][bold #818cf8]███╗   ███╗[/bold #818cf8]\n"
    "[bold #00ffff]  ╚███╗  [/bold #00ffff]         [bold #00ffff]██╔══██╗[/bold #00ffff][bold #22d3ee]██╔════╝[/bold #22d3ee][bold #38bdf8]██╔══██╗[/bold #38bdf8][bold #60a5fa]██║[/bold #60a5fa][bold #818cf8]████╗ ████║[/bold #818cf8]\n"
    "[bold #00ffff]   ╚███╗ [/bold #00ffff]         [bold #00ffff]██║  ██║[/bold #00ffff][bold #22d3ee]█████╗  [/bold #22d3ee][bold #38bdf8]██████╔╝[/bold #38bdf8][bold #60a5fa]██║[/bold #60a5fa][bold #818cf8]██╔████╔██║[/bold #818cf8]\n"
    "[bold #00ffff]   ╔███╝ [/bold #00ffff]         [bold #00ffff]██║  ██║[/bold #00ffff][bold #22d3ee]██╔══╝  [/bold #22d3ee][bold #38bdf8]██╔══██╗[/bold #38bdf8][bold #60a5fa]██║[/bold #60a5fa][bold #818cf8]██║╚██╔╝██║[/bold #818cf8]\n"
    "[bold #00ffff]  ╔███╝  [/bold #00ffff] [bold #ffd700]██████╗[/bold #ffd700] [bold #00ffff]██████╔╝[/bold #00ffff][bold #22d3ee]███████╗[/bold #22d3ee][bold #38bdf8]██████╔╝[/bold #38bdf8][bold #60a5fa]██║[/bold #60a5fa][bold #818cf8]██║ ╚═╝ ██║[/bold #818cf8]\n"
    "[bold #00ffff]  ╚═══╝  [/bold #00ffff] [bold #ffd700]╚═════╝[/bold #ffd700] [bold #00ffff]╚═════╝ [/bold #00ffff][bold #22d3ee]╚══════╝[/bold #22d3ee][bold #38bdf8]╚═════╝ [/bold #38bdf8][bold #60a5fa]╚═╝[/bold #60a5fa][bold #818cf8]╚═╝     ╚═╝[/bold #818cf8]\n"
    "[dim #64748b]  -------------------------------------------------------------[/dim #64748b]\n"
    f"[bold #94a3b8]   Declarative BIM Compiler v{__version__}  |  Developer: Prida Takon[/bold #94a3b8]\n"
    "[bold #64748b]   Building-as-Code for AI Agents and Humans (@PRIDA-TAKON)[/bold #64748b]\n"
)


# --- Command Handlers ---

def cmd_init(args):
    project_dir = Path(args.name)
    if project_dir.exists():
        console.print(f"[bold red]Error:[/bold red] Directory '{args.name}' already exists.")
        sys.exit(1)

    project_dir.mkdir(parents=True, exist_ok=True)
    (project_dir / "assets").mkdir(exist_ok=True)
    (project_dir / "tests").mkdir(exist_ok=True)
    console.print(
        f"[bold green]Success:[/bold green] Initialized new debim project in [cyan]{args.name}[/cyan]"
    )


def cmd_validate(args):
    manifest_target = args.manifest or getattr(args, "positional_manifest", None)
    manifest = resolve_default_manifest_path(manifest_target)
    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)
    console.print(f"[bold green]Validating:[/bold green] {manifest}")
    has_diagnostic_warnings = False
    try:
        manifest_obj = load_manifest(manifest)
        console.print("[bold green]Pydantic Schema Validation passed.[/bold green]")

        # 1. Topological continuity check
        resolved_manifest = resolve_manifest(manifest_obj)
        topo_warnings = check_topological_continuity(resolved_manifest)
        if topo_warnings:
            has_diagnostic_warnings = True
            topo_table = Table(
                title="Topological Continuity Diagnostic Warnings",
                show_header=True,
                header_style="bold yellow",
            )
            topo_table.add_column("Type", style="bold red")
            topo_table.add_column("Tag", style="bold yellow")
            topo_table.add_column("Class", style="white")
            topo_table.add_column("Position", style="cyan")
            topo_table.add_column("Message", style="dim")

            for tw in topo_warnings:
                pos_str = f"({tw['position'][0]:.2f}, {tw['position'][1]:.2f}, {tw['position'][2]:.2f})"
                topo_table.add_row(
                    tw["type"],
                    tw["tag"],
                    tw["class"],
                    pos_str,
                    tw["message"],
                )
            console.print(topo_table)
        else:
            console.print("[bold green]Topological Continuity Validation passed (no floating columns or unsupported beams).[/bold green]")

        # 2. MEP vs Structural AABB Clash Detection
        mep_clashes = detect_mep_structural_clashes(resolved_manifest)
        if mep_clashes:
            has_diagnostic_warnings = True
            clash_table = Table(
                title="MEP vs Structural AABB Clash Detection Summary",
                show_header=True,
                header_style="bold yellow",
            )
            clash_table.add_column("MEP Tag", style="bold yellow")
            clash_table.add_column("MEP Class", style="cyan")
            clash_table.add_column("Structural Tag", style="bold magenta")
            clash_table.add_column("Structural Class", style="white")
            clash_table.add_column("Diagnostic Message", style="dim")

            for c in mep_clashes:
                clash_table.add_row(
                    c["mep_tag"],
                    c["mep_class"],
                    c["struct_tag"],
                    c["struct_class"],
                    c["message"],
                )
            console.print(clash_table)
        else:
            console.print("[bold green]AABB Clash Detection passed (no MEP-structure intersections).[/bold green]")

        # 3. bSDD Property Sets validation
        if getattr(args, "bsdd", True):
            pset_results = validate_manifest_psets(manifest_obj)
            if pset_results:
                table = Table(
                    title="buildingSMART bSDD Pset Validation Summary",
                    show_header=True,
                    header_style="bold cyan",
                )
                table.add_column("Tag", style="bold yellow")
                table.add_column("Property Set", style="magenta")
                table.add_column("Property", style="white")
                table.add_column("Expected Type", style="cyan")
                table.add_column("Status", justify="center")
                table.add_column("Details", style="dim")

                has_pset_errors = False
                for r in pset_results:
                    if r.status in (PropertyValidationStatus.TYPE_MISMATCH, PropertyValidationStatus.VALUE_CONSTRAINT_VIOLATION):
                        has_pset_errors = True
                        status_str = f"[bold red]{r.status.value}[/bold red]"
                    elif r.status in (PropertyValidationStatus.NON_STANDARD_PROPERTY, PropertyValidationStatus.NON_STANDARD_PSET):
                        status_str = f"[yellow]{r.status.value}[/yellow]"
                    else:
                        status_str = f"[green]{r.status.value}[/green]"

                    table.add_row(
                        r.element_tag,
                        r.pset_name,
                        r.property_name,
                        r.expected_type or "-",
                        status_str,
                        r.message,
                    )

                console.print(table)

                if getattr(args, "strict_psets", False) and (has_pset_errors or any(r.status != PropertyValidationStatus.VALID for r in pset_results)):
                    console.print("[bold red]Strict Pset Validation Failed! Non-compliant Property Sets detected.[/bold red]")
                    sys.exit(1)

        if has_diagnostic_warnings:
            console.print("[bold yellow]Validation completed with diagnostic warnings.[/bold yellow]")
            sys.exit(1)
        else:
            console.print("[bold green]Validation completed successfully with 0 warnings or errors.[/bold green]")

    except (ValidationError, ValueError) as e:
        console.print(f"[bold red]Validation Error:[/bold red]\n{e}")
        sys.exit(1)
    except SystemExit:
        raise
    except Exception as e:
        console.print(f"[bold red]Validation Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_schema(args):
    import json

    output_path = Path(args.output) if args.output else None
    schema_file = output_path or (Path("debim.schema.json") if args.vscode else None)
    schema = export_json_schema(output_path=schema_file, indent=args.indent)

    if schema_file:
        defs_count = len(schema.get("$defs", {}))
        console.print(
            f"[bold green]Success:[/bold green] Exported debim JSON Schema ({defs_count} definitions) to [cyan]{schema_file}[/cyan]"
        )
    else:
        sys.stdout.write(json.dumps(schema, indent=args.indent, ensure_ascii=False) + "\n")

    if args.vscode:
        vscode_dir = Path(".vscode")
        vscode_dir.mkdir(parents=True, exist_ok=True)
        settings_file = vscode_dir / "settings.json"
        settings_data = {}
        if settings_file.exists():
            try:
                with open(settings_file, "r", encoding="utf-8") as f:
                    settings_data = json.load(f)
            except Exception:
                settings_data = {}

        associations = settings_data.setdefault("files.associations", {})
        associations["*.debim"] = "yaml"
        yaml_schemas = settings_data.setdefault("yaml.schemas", {})
        schema_rel_path = f"./{schema_file.name if schema_file else 'debim.schema.json'}"
        patterns = ["project.yaml", "**/project.yaml", "*.debim.yaml", "**/*.debim.yaml", "*.debim", "**/*.debim"]
        yaml_schemas[schema_rel_path] = patterns

        with open(settings_file, "w", encoding="utf-8") as f:
            json.dump(settings_data, f, indent=2, ensure_ascii=False)
            f.write("\n")

        console.print(
            f"[bold green]Success:[/bold green] Configured [cyan]{settings_file}[/cyan] with schema mapping {patterns}."
        )


def cmd_split(args):
    manifest = resolve_default_manifest_path(args.manifest)
    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    output_dir = Path(args.output_dir) if args.output_dir else None
    try:
        summary_info = split_manifest(
            manifest_path=manifest,
            output_dir=output_dir,
            by=args.by,
            backup=args.backup,
            dry_run=args.dry_run,
        )
    except Exception as e:
        console.print(f"[bold red]Error during split:[/bold red] {e}")
        sys.exit(1)

    title = "[bold cyan]debim Split (Preview / Dry Run)[/bold cyan]" if args.dry_run else "[bold green]debim Split Complete[/bold green]"
    table = Table(title=title, show_header=True, header_style="bold cyan")
    table.add_column("Module File", style="yellow")
    table.add_column("Elements", justify="right", style="green")

    for rel_path, count in summary_info["modules"].items():
        table.add_row(rel_path, str(count))

    console.print(table)
    total_elements = summary_info["total_elements"]
    num_modules = len(summary_info["modules"])

    if args.dry_run:
        console.print(
            Panel(
                f"[bold yellow]Dry-run preview:[/bold yellow] {total_elements} elements across {num_modules} modules.\n"
                f"Strategy: [cyan]{args.by}[/cyan]\n"
                "Run without [cyan]--dry-run[/cyan] to write modular manifest files to disk.",
                style="yellow",
            )
        )
    else:
        console.print(
            Panel(
                f"[bold green]Successfully decomposed[/bold green] [cyan]{manifest}[/cyan] into [yellow]{num_modules}[/yellow] modular files!\n"
                f"Total Elements: [green]{total_elements}[/green] | Strategy: [cyan]{args.by}[/cyan]\n"
                f"Root Manifest Updated: [cyan]{summary_info['root_file']}[/cyan] (with includes: [\"models/**/*.yaml\"])",
                style="green",
            )
        )


def cmd_bundle(args):
    manifest = resolve_default_manifest_path(args.manifest)
    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    output = Path(args.output) if args.output else None
    try:
        bundled = bundle_manifest(manifest_path=manifest, output_path=output)
    except Exception as e:
        console.print(f"[bold red]Error during bundle:[/bold red] {e}")
        sys.exit(1)

    elem_count = len(bundled.get("elements", []))
    storey_count = len(bundled.get("spatial_structure", {}).get("storeys", []))
    mat_count = len(bundled.get("materials", []))

    if output is not None:
        console.print(
            Panel(
                f"[bold green]Successfully bundled modular manifest into:[/bold green] [cyan]{output}[/cyan]\n"
                f"Aggregated: [green]{elem_count}[/green] elements, [green]{storey_count}[/green] storeys, [green]{mat_count}[/green] materials.",
                title="[bold green]debim Bundle Complete[/bold green]",
                style="green",
            )
        )
    else:
        yaml_str = yaml.safe_dump(bundled, sort_keys=False, allow_unicode=True)
        print(yaml_str, end="")


def cmd_test(args):
    import subprocess

    test_dir = Path(args.test_dir)
    console.print(
        Panel(
            f"[bold cyan]Automated Building Compliance Engine & Pytest Checker[/bold cyan]\n"
            f"[dim]Running compliance rule suite in:[/dim] [yellow]{test_dir}[/yellow]",
            title="[bold green]debim Compliance Engine[/bold green]",
        )
    )

    cmd = [sys.executable, "-m", "pytest", str(test_dir), "-rA"]
    if args.verbose:
        cmd.append("-v")

    res = subprocess.run(cmd, capture_output=True, text=True)

    if res.stdout:
        console.print(res.stdout)
    if res.stderr:
        console.print(f"[dim]{res.stderr}[/dim]")

    if res.returncode == 0:
        console.print(
            Panel(
                "[bold green]ALL COMPLIANCE RULES PASSED ACCORDING TO BUILDING CODES![/bold green]",
                style="green",
            )
        )
    else:
        console.print(
            Panel(
                "[bold red]COMPLIANCE CHECKS FAILED! PLEASE REVIEW REGULATORY VIOLATIONS ABOVE.[/bold red]",
                style="red",
            )
        )

    sys.exit(res.returncode)


def cmd_summary(args):
    manifest = resolve_default_manifest_path(args.manifest)
    prices = Path(args.prices)
    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    try:
        manifest_obj = load_manifest(manifest)
        resolved = resolve_manifest(manifest_obj)
        project_qto = calculate_qto(resolved)

        proj = manifest_obj.project
        console.print(
            Panel(
                f"[bold cyan]Project Name:[/bold cyan] {proj.name}\n"
                f"[bold cyan]Project ID:[/bold cyan] {proj.id}\n"
                f"[bold cyan]Units:[/bold cyan] length={proj.units.length}, area={proj.units.area}, volume={proj.units.volume}",
                title="[bold green]Project Overview[/bold green]",
            )
        )

        axes_x = list(manifest_obj.grids.axes_x.values())
        axes_y = list(manifest_obj.grids.axes_y.values())
        min_x, max_x = (min(axes_x), max(axes_x)) if axes_x else (0.0, 0.0)
        min_y, max_y = (min(axes_y), max(axes_y)) if axes_y else (0.0, 0.0)
        span_x = max_x - min_x
        span_y = max_y - min_y
        bounding_area = span_x * span_y

        site_table = Table(title="Site Bounds & Spatial Structure", show_header=True, header_style="bold cyan")
        site_table.add_column("Property", style="bold yellow")
        site_table.add_column("Value", style="white")

        site_table.add_row("X Range", f"{min_x:.2f} m to {max_x:.2f} m (Span: {span_x:.2f} m)")
        site_table.add_row("Y Range", f"{min_y:.2f} m to {max_y:.2f} m (Span: {span_y:.2f} m)")
        site_table.add_row("Bounding Area", f"{bounding_area:.2f} m2")

        storeys_str = ", ".join([f"{s.name} (+{s.elevation:.2f}m)" for s in manifest_obj.spatial_structure.storeys])
        site_table.add_row("Storeys", storeys_str)

        console.print(site_table)

        class_counts = {}
        for elem in manifest_obj.elements:
            cls = elem.class_
            class_counts[cls] = class_counts.get(cls, 0) + 1

        door_count = len(resolved.doors)
        window_count = len(resolved.windows)
        if door_count > 0:
            class_counts["IfcDoor"] = door_count
        if window_count > 0:
            class_counts["IfcWindow"] = window_count

        elem_table = Table(title="Element Statistics Breakdown", show_header=True, header_style="bold cyan")
        elem_table.add_column("IFC Class", style="bold green")
        elem_table.add_column("Count", justify="right", style="bold yellow")

        for cls, count in sorted(class_counts.items()):
            elem_table.add_row(cls, str(count))

        console.print(elem_table)

        cost_str = "N/A (prices.json not found)"
        if prices.exists():
            try:
                price_catalog = load_price_catalog(prices)
                estimate = estimate_cost(project_qto, price_catalog, manifest_obj)
                cost_str = f"{estimate.grand_total:,.2f} {estimate.currency}"
            except Exception:
                pass

        qto_summary = (
            f"[bold cyan]Total Concrete Volume:[/bold cyan] {project_qto.total_concrete_volume:.3f} m3\n"
            f"[bold cyan]Total Formwork Area:[/bold cyan] {project_qto.total_formwork_area:.3f} m2\n"
            f"[bold cyan]Total Rebar Weight:[/bold cyan] {project_qto.total_rebar_weight:.3f} kg\n"
            f"[bold green]Estimated Budget:[/bold green] {cost_str}"
        )
        console.print(Panel(qto_summary, title="[bold green]QTO & Cost Snapshot[/bold green]"))

    except Exception as e:
        console.print(f"[bold red]Summary Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_qto(args):
    manifest = resolve_default_manifest_path(args.manifest)
    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    console.print(f"[bold blue]Calculating QTO for:[/bold blue] {manifest}")
    try:
        manifest_obj = load_manifest(manifest)
        project_qto = calculate_qto(manifest_obj)

        table = Table(
            title="Quantitative Take-Off (QTO) Summary",
            show_header=True,
            header_style="bold cyan",
        )
        table.add_column("Tag", style="bold yellow")
        table.add_column("Class", style="green")
        table.add_column("Material", style="magenta")
        table.add_column("Concrete Vol (m3)", justify="right")
        table.add_column("Formwork Area (m2)", justify="right")
        table.add_column("Rebar Weight (kg)", justify="right")

        for eqto in project_qto.elements:
            table.add_row(
                eqto.tag,
                eqto.element_class,
                eqto.material or "-",
                f"{eqto.concrete_volume:.3f}",
                f"{eqto.formwork_area:.3f}",
                f"{eqto.total_rebar_weight:.3f}",
            )

        console.print(table)

        summary_text = (
            f"[bold cyan]Total Concrete Volume:[/bold cyan] {project_qto.total_concrete_volume:.3f} m3\n"
            f"[bold cyan]Total Formwork Area:[/bold cyan] {project_qto.total_formwork_area:.3f} m2\n"
            f"[bold cyan]Total Rebar Weight:[/bold cyan] {project_qto.total_rebar_weight:.3f} kg"
        )
        if project_qto.total_excavation_volume > 0 or project_qto.total_lean_concrete_volume > 0 or project_qto.total_sand_bedding_volume > 0:
            summary_text += (
                f"\n[bold magenta]--- Substructure & Earthwork ---[/bold magenta]\n"
                f"[bold cyan]Excavation (งานดินขุดฐานราก):[/bold cyan] {project_qto.total_excavation_volume:.2f} m3\n"
                f"[bold cyan]Lean Concrete (คอนกรีตหยาบ):[/bold cyan] {project_qto.total_lean_concrete_volume:.2f} m3\n"
                f"[bold cyan]Sand Bedding (ทรายหยาบรองฐานราก/พื้น):[/bold cyan] {project_qto.total_sand_bedding_volume:.2f} m3"
            )
        if project_qto.total_ceiling_gypsum_area > 0 or project_qto.total_ceiling_tbar_area > 0 or project_qto.total_ceiling_eaves_area > 0:
            summary_text += (
                f"\n[bold magenta]--- Ceilings (งานฝ้าเพดาน) ---[/bold magenta]\n"
                f"[bold cyan]Gypsum Ceiling (ฝ้ายิปซั่มฉาบเรียบ):[/bold cyan] {project_qto.total_ceiling_gypsum_area:.2f} m2\n"
                f"[bold cyan]T-Bar Ceiling (ฝ้าทีบาร์ทนชื้นห้องน้ำ):[/bold cyan] {project_qto.total_ceiling_tbar_area:.2f} m2\n"
                f"[bold cyan]Eaves Ceiling (ฝ้าชายคาระบายอากาศ):[/bold cyan] {project_qto.total_ceiling_eaves_area:.2f} m2"
            )
        if project_qto.total_floor_tile_area > 0 or project_qto.total_floor_polish_area > 0 or project_qto.total_skirting_length > 0:
            summary_text += (
                f"\n[bold magenta]--- Flooring & Finishes (งานปูพื้นและตกแต่งผิว) ---[/bold magenta]\n"
                f"[bold cyan]Ceramic Floor Tiles (ปูกระเบื้องเซรามิค):[/bold cyan] {project_qto.total_floor_tile_area:.2f} m2\n"
                f"[bold cyan]Polished Concrete (คอนกรีตขัดเรียบ/ขัดมัน):[/bold cyan] {project_qto.total_floor_polish_area:.2f} m2\n"
                f"[bold cyan]Skirting Board (บัวเชิงผนัง):[/bold cyan] {project_qto.total_skirting_length:.2f} m"
            )
        if project_qto.total_structural_steel_weight > 0 or project_qto.total_painting_area > 0 or project_qto.total_timber_volume > 0:
            summary_text += (
                f"\n[bold magenta]--- Structural Steel & Timber (งานโครงสร้างเหล็กและไม้) ---[/bold magenta]\n"
            )
            if project_qto.total_structural_steel_weight > 0:
                summary_text += f"[bold cyan]Structural Steel Weight (เหล็กรูปพรรณ):[/bold cyan] {project_qto.total_structural_steel_weight:.2f} kg ({project_qto.total_structural_steel_weight/1000.0:.3f} tons)\n"
            if project_qto.total_painting_area > 0:
                summary_text += f"[bold cyan]Painting Area (พื้นที่ทาสีจริง):[/bold cyan] {project_qto.total_painting_area:.2f} m2\n"
            if project_qto.total_weld_touchup_area > 0:
                summary_text += f"[bold cyan]Weld Touch-Up Area (พื้นที่สีกันสนิมรอยเชื่อม 10%):[/bold cyan] {project_qto.total_weld_touchup_area:.2f} m2\n"
            if project_qto.total_timber_volume > 0:
                summary_text += f"[bold cyan]Timber Volume (ปริมาตรไม้โครงสร้าง):[/bold cyan] {project_qto.total_timber_volume:.3f} m3\n"
        if project_qto.total_roof_covering_area > 0:
            summary_text += (
                f"\n[bold cyan]Roof Covering Tiles:[/bold cyan] {project_qto.total_roof_covering_area:.2f} m2\n"
                f"[bold cyan]Roof Structural Steel:[/bold cyan] {project_qto.total_roof_steel_weight:.2f} kg\n"
                f"[bold cyan]Roof Ridge/Hip Caps:[/bold cyan] {project_qto.total_roof_ridge_length + project_qto.total_roof_hip_length:.2f} m\n"
                f"[bold cyan]Roof Eaves Perimeter:[/bold cyan] {project_qto.total_roof_eaves_length:.2f} m"
            )
        total_pipe = (
            project_qto.total_cold_water_pipe_length
            + project_qto.total_soil_pipe_length
            + project_qto.total_waste_pipe_length
            + project_qto.total_vent_pipe_length
            + project_qto.total_drainage_pipe_length
            + project_qto.total_refrigerant_pipe_length
            + project_qto.total_condensate_pipe_length
        )
        if (
            total_pipe > 0
            or project_qto.total_sanitary_terminals_count > 0
            or project_qto.total_conduit_length > 0
            or project_qto.total_duct_length > 0
            or project_qto.total_air_terminals_count > 0
            or project_qto.total_unitary_equipment_count > 0
        ):
            summary_text += (
                f"\n[bold magenta]--- MEP System Totals ---[/bold magenta]\n"
                f"[bold cyan]Cold Water Pipe (ท่อน้ำดี):[/bold cyan] {project_qto.total_cold_water_pipe_length:.2f} m\n"
                f"[bold cyan]Soil & Waste Pipe (ท่อโสโครก/น้ำทิ้ง):[/bold cyan] {project_qto.total_soil_pipe_length + project_qto.total_waste_pipe_length:.2f} m\n"
                f"[bold cyan]Vent Pipe (ท่อระบายอากาศ):[/bold cyan] {project_qto.total_vent_pipe_length:.2f} m\n"
                f"[bold cyan]Pipe Fittings & Valves (ข้อต่อท่อ):[/bold cyan] {project_qto.total_pipe_fittings_count} items\n"
                f"[bold cyan]Sanitary Fixtures (สุขภัณฑ์/อุปกรณ์):[/bold cyan] {project_qto.total_sanitary_terminals_count} sets\n"
                f"[bold cyan]Electrical Conduits (ท่อร้อยสายไฟ):[/bold cyan] {project_qto.total_conduit_length:.2f} m\n"
                f"[bold cyan]Distribution Boards (ตู้ควบคุมไฟ):[/bold cyan] {project_qto.total_distribution_boards_count} sets\n"
                f"[bold cyan]Lighting Fixtures (ดวงโคม):[/bold cyan] {project_qto.total_lighting_fixtures_count} sets\n"
                f"[bold cyan]Switches & Outlets (สวิตช์/เต้ารับ):[/bold cyan] {project_qto.total_switches_count + project_qto.total_outlets_count} sets\n"
                f"[bold cyan]Refrigerant Pipe (ท่อน้ำยาแอร์):[/bold cyan] {project_qto.total_refrigerant_pipe_length:.2f} m\n"
                f"[bold cyan]HVAC Ducts (ท่อลมระบายอากาศ):[/bold cyan] {project_qto.total_duct_length:.2f} m\n"
                f"[bold cyan]Air Conditioners (เครื่องปรับอากาศ):[/bold cyan] {project_qto.total_unitary_equipment_count} sets\n"
                f"[bold cyan]Ventilation Fans & Hoods (พัดลมดูด/ฮูด):[/bold cyan] {project_qto.total_air_terminals_count} sets"
            )
        console.print(Panel(summary_text, title="[bold green]QTO Totals[/bold green]"))

    except Exception as e:
        console.print(f"[bold red]QTO Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_cost(args):
    manifest = resolve_default_manifest_path(args.manifest)
    prices = Path(args.prices)
    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    prices_str = str(prices)
    if prices_str != "-" and not prices.exists():
        console.print(f"[bold red]Error:[/bold red] Price catalog '{prices}' not found.")
        sys.exit(1)

    console.print(
        f"[bold blue]Calculating project cost for:[/bold blue] {manifest} [blue]using[/blue] {prices}"
    )
    try:
        manifest_obj = load_manifest(manifest)
        project_qto = calculate_qto(manifest_obj)
        price_catalog = load_price_catalog(prices)

        estimate = estimate_cost(project_qto, price_catalog, manifest_obj)

        table = Table(
            title=f"Cost Estimate Summary ({estimate.currency})",
            show_header=True,
            header_style="bold cyan",
        )
        table.add_column("Code", style="bold yellow")
        table.add_column("Description", style="white")
        table.add_column("Unit", justify="center")
        table.add_column("Quantity", justify="right")
        table.add_column("Unit Mat", justify="right")
        table.add_column("Unit Labor", justify="right")
        table.add_column("Total Mat", justify="right")
        table.add_column("Total Labor", justify="right")
        table.add_column("Total Amount", justify="right", style="bold green")

        for item in estimate.line_items:
            table.add_row(
                item.code,
                item.name,
                item.unit,
                f"{item.quantity:.3f}",
                f"{item.unit_material_cost:,.2f}",
                f"{item.unit_labor_cost:,.2f}",
                f"{item.total_material_cost:,.2f}",
                f"{item.total_labor_cost:,.2f}",
                f"{item.total_amount:,.2f}",
            )

        console.print(table)

        summary_text = (
            f"[bold cyan]Total Material Cost:[/bold cyan] {estimate.total_material_cost:,.2f} {estimate.currency}\n"
            f"[bold cyan]Total Labor Cost:[/bold cyan] {estimate.total_labor_cost:,.2f} {estimate.currency}\n"
            f"[bold green]Grand Total Cost:[/bold green] {estimate.grand_total:,.2f} {estimate.currency}"
        )
        console.print(
            Panel(summary_text, title="[bold green]Project Budget Summary[/bold green]")
        )

        if args.output:
            csv_path = estimate.export_csv(Path(args.output))
            console.print(
                f"[bold green]Exported BOQ CSV to:[/bold green] [cyan]{csv_path}[/cyan]"
            )

    except Exception as e:
        console.print(f"[bold red]Cost Estimation Error:[/bold red]\n{e}")
        sys.exit(1)


# Thai Provincial Price Index Multiplier Lookup Table (Baseline: Bangkok / Central = 1.00)
PROVINCIAL_PRICE_INDEX: Dict[str, float] = {
    "BANGKOK": 1.00,
    "กรุงเทพมหานคร": 1.00,
    "NONTHABURI": 1.00,
    "PATHUM THANI": 1.00,
    "CHIANG MAI": 1.05,
    "เชียงใหม่": 1.05,
    "PHUKET": 1.15,
    "ภูเก็ต": 1.15,
    "CHON BURI": 1.02,
    "ชลบุรี": 1.02,
    "KHON KAEN": 1.06,
    "ขอนแก่น": 1.06,
    "SONGKHLA": 1.08,
    "สงขลา": 1.08,
}


def cmd_boq(args):
    manifest_target = args.manifest or args.positional_manifest
    manifest = resolve_default_manifest_path(manifest_target)
    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    rules_mode = (args.rules or "cg-v480").lower()
    province_name = args.province.strip() if args.province else None
    province_mult = 1.00
    if province_name:
        prov_key = province_name.upper()
        province_mult = PROVINCIAL_PRICE_INDEX.get(prov_key, PROVINCIAL_PRICE_INDEX.get(province_name, 1.05))

    console.print(
        f"[bold blue]Generating Bill of Quantities (BOQ) for:[/bold blue] {manifest} "
        f"[dim](Rules: {rules_mode.upper()}, Province: {province_name or 'Bangkok Standard'}{f' [x{province_mult:.2f}]' if province_name else ''})[/dim]"
    )

    try:
        manifest_obj = load_manifest(manifest)
        project_qto = calculate_qto(manifest_obj)

        # Generate price catalog fallback if prices file not specified or missing
        prices_path = Path(args.prices) if args.prices else None
        if prices_path and prices_path.exists():
            price_catalog = load_price_catalog(prices_path)
        else:
            import tempfile
            with tempfile.NamedTemporaryFile(suffix=".yaml", delete=False) as tmp:
                tmp_path = Path(tmp.name)
            tmpl_path = generate_cost_template(manifest_obj, output_path=tmp_path)
            price_catalog = load_price_catalog(tmpl_path)
            try:
                tmp_path.unlink()
            except Exception:
                pass
            # Default realistic benchmark prices if template was newly generated
            for item in price_catalog.items.values():
                if item.unit.lower() == "m3":
                    item.material_cost = item.material_cost or 2400.0
                    item.labor_cost = item.labor_cost or 450.0
                elif item.unit.lower() == "m2":
                    item.material_cost = item.material_cost or 350.0
                    item.labor_cost = item.labor_cost or 120.0
                elif item.unit.lower() == "kg":
                    item.material_cost = item.material_cost or 32.0
                    item.labor_cost = item.labor_cost or 5.0
                elif item.unit.lower() == "m":
                    item.material_cost = item.material_cost or 180.0
                    item.labor_cost = item.labor_cost or 60.0
                else:
                    item.material_cost = item.material_cost or 1500.0
                    item.labor_cost = item.labor_cost or 500.0

        # Apply Provincial Price Index Multiplier
        if province_mult != 1.00:
            for item in price_catalog.items.values():
                item.material_cost = round(item.material_cost * province_mult, 2)
                item.labor_cost = round(item.labor_cost * province_mult, 2)

        estimate = estimate_cost(project_qto, price_catalog, manifest_obj)
        proj_name = manifest_obj.project.name or "โครงการก่อสร้าง"

        # Generate Thai Standard BOQ Sheets (ปร.4, ปร.4(พ), ปร.5)
        pr4 = estimate.generate_pr4(project_name=proj_name)
        pr4_p = estimate.generate_pr4_p(project_qto, price_catalog, project_name=proj_name)
        pr5 = estimate.generate_pr5(project_name=proj_name)

        # 1. Print ANSI formatted summary table for แบบ ปร.4
        pr4_table = Table(
            title=f"แบบ ปร.4 - รายการปริมาณงานและราคา ({pr4.project_name})",
            show_header=True,
            header_style="bold cyan",
        )
        pr4_table.add_column("ลำดับ", justify="center", style="dim")
        pr4_table.add_column("รายการ (Description)", style="white")
        pr4_table.add_column("หน่วย", justify="center", style="cyan")
        pr4_table.add_column("จำนวน", justify="right", style="yellow")
        pr4_table.add_column("ค่าวัสดุรวม", justify="right")
        pr4_table.add_column("ค่าแรงงานรวม", justify="right")
        pr4_table.add_column("รวมเป็นเงิน (THB)", justify="right", style="bold green")

        for item in pr4.items:
            pr4_table.add_row(
                str(item.item_no),
                item.description,
                item.unit,
                f"{item.quantity:,.2f}",
                f"{item.material_total_cost:,.2f}",
                f"{item.labor_total_cost:,.2f}",
                f"{item.total_amount:,.2f}",
            )
        console.print(pr4_table)

        # 2. Print Summary for แบบ ปร.4(พ)
        if pr4_p.items:
            pr4p_table = Table(
                title=f"แบบ ปร.4(พ) - สรุปตารางเหล็กเสริมและไม้แบบ ({pr4_p.project_name})",
                show_header=True,
                header_style="bold magenta",
            )
            pr4p_table.add_column("หมวดงาน", style="cyan")
            pr4p_table.add_column("ประเภท / ขนาด", style="white")
            pr4p_table.add_column("ปริมาณ", justify="right", style="yellow")
            pr4p_table.add_column("หน่วย", justify="center")
            pr4p_table.add_column("รวมเป็นเงิน (THB)", justify="right", style="bold green")

            for item in pr4_p.items:
                pr4p_table.add_row(
                    item.category,
                    item.bar_type_or_type,
                    f"{item.quantity:,.2f}",
                    item.unit,
                    f"{item.total_amount:,.2f}",
                )
            console.print(pr4p_table)

        # 3. Print แบบ ปร.5 Summary Panel (Direct Cost + Factor F)
        summary_panel_text = (
            f"[bold cyan]โครงการ:[/bold cyan] {pr5.project_name}\n"
            f"[bold cyan]ราคาทุนทางตรง (Direct Cost):[/bold cyan] {pr5.direct_cost:,.2f} THB\n"
            f"[bold yellow]ตัวคูณ Factor F (กรมบัญชีกลาง ว480):[/bold yellow] [bold green]{pr5.factor_f:.4f}[/bold green]\n"
            f"[bold green]ราคากลางรวมค่าก่อสร้าง (Grand Total Construction Cost):[/bold green] [bold green]{pr5.total_construction_cost:,.2f} THB[/bold green]\n"
            f"[dim](รวมภาษีมูลค่าเพิ่ม VAT 7% และ Factor F ตามประกาศกรมบัญชีกลาง ว480)[/dim]"
        )
        if province_name:
            summary_panel_text += f"\n[bold magenta]ดัชนีราคาจังหวัด ({province_name}):[/bold magenta] x{province_mult:.2f}"

        console.print(Panel(summary_panel_text, title="[bold green]แบบ ปร.5 - สรุปราคากลางค่าก่อสร้าง[/bold green]", style="green"))

        # 4. Handle --export-excel
        if args.export_excel:
            out_excel = Path(args.export_excel)
            out_excel.parent.mkdir(parents=True, exist_ok=True)

            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

                wb = openpyxl.Workbook()

                # Sheet 1: ปร.5 Summary
                ws_pr5 = wb.active
                ws_pr5.title = "ปร.5 สรุปราคากลาง"
                ws_pr5.append(["แบบ ปร.5 - แบบสรุปค่าก่อสร้าง"])
                ws_pr5.append([f"โครงการ: {pr5.project_name}"])
                ws_pr5.append([])
                ws_pr5.append(["กลุ่มงาน / รายการ", "ค่างานต้นทุน (บาท)", "Factor F", "รวมค่าก่อสร้าง (บาท)"])
                ws_pr5.append(["1. กลุ่มงานอาคาร (Building Works)", pr5.direct_cost, pr5.factor_f, pr5.total_construction_cost])
                ws_pr5.append(["รวมราคากลางค่าก่อสร้างทั้งหมด", pr5.direct_cost, pr5.factor_f, pr5.total_construction_cost])

                # Sheet 2: ปร.4 Detail
                ws_pr4 = wb.create_sheet(title="ปร.4 รายการปริมาณงาน")
                ws_pr4.append(["แบบ ปร.4 - รายการปริมาณงานและราคา"])
                ws_pr4.append(["ลำดับ", "รายการ", "หน่วย", "จำนวน", "ค่าวัสดุต่อหน่วย", "ค่าวัสดุรวม", "ค่าแรงงานต่อหน่วย", "ค่าแรงงานรวม", "รวมเป็นเงิน"])
                for item in pr4.items:
                    ws_pr4.append([
                        item.item_no,
                        item.description,
                        item.unit,
                        item.quantity,
                        item.material_unit_cost,
                        item.material_total_cost,
                        item.labor_unit_cost,
                        item.labor_total_cost,
                        item.total_amount,
                    ])
                ws_pr4.append(["รวม", "รวมค่างานต้นทุนทางตรง", "", "", "", pr4.total_material_cost, "", pr4.total_labor_cost, pr4.total_direct_cost])

                # Sheet 3: ปร.4(พ) Rebar & Formwork
                ws_pr4p = wb.create_sheet(title="ปร.4(พ) เหล็กเสริม-ไม้แบบ")
                ws_pr4p.append(["แบบ ปร.4(พ) - แบบสรุปพิกัดตารางเหล็กเสริมและไม้แบบ"])
                ws_pr4p.append(["ลำดับ", "หมวดงาน", "ประเภท/ขนาด", "หน่วย", "ปริมาณ", "ค่าวัสดุต่อหน่วย", "ค่าแรงงานต่อหน่วย", "รวมเป็นเงิน"])
                for item in pr4_p.items:
                    ws_pr4p.append([
                        item.item_no,
                        item.category,
                        item.bar_type_or_type,
                        item.unit,
                        item.quantity,
                        item.unit_material_cost,
                        item.unit_labor_cost,
                        item.total_amount,
                    ])
                ws_pr4p.append(["รวม", "รวมตารางเหล็กเสริมและไม้แบบ", "", "", "", "", "", pr4_p.total_amount])

                wb.save(out_excel)
                console.print(f"[bold green]Exported BOQ Excel spreadsheet to:[/bold green] [cyan]{out_excel}[/cyan]")

            except ImportError:
                # Fallback to CSV export if openpyxl is not installed
                csv_out = out_excel.with_suffix(".csv")
                estimate.export_csv(csv_out)
                console.print(
                    f"[bold yellow]openpyxl not installed. Exported BOQ breakdown to CSV:[/bold yellow] [cyan]{csv_out}[/cyan]"
                )

    except Exception as e:
        console.print(f"[bold red]BOQ Generation Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_cost_template(args):
    manifest = resolve_default_manifest_path(args.manifest)
    output = Path(args.output) if args.output else None

    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    if output is None:
        if args.modular:
            output = Path("prices/catalog.yaml")
        else:
            ext = "json" if args.format.lower() == "json" else "yaml"
            output = Path(f"prices.template.{ext}")

    console.print(
        f"[bold green]Generating cost template from:[/bold green] {manifest} -> [cyan]{output}[/cyan]"
    )
    try:
        out_path = generate_cost_template(
            manifest=manifest,
            output_path=output,
            format=args.format,
            modular=args.modular,
        )
        console.print(
            Panel(
                f"[bold green]Cost Catalog Template Generated Successfully![/bold green]\n"
                f"[bold cyan]Output File:[/bold cyan] {out_path}\n"
                f"[bold cyan]Modular Mode:[/bold cyan] {args.modular}\n"
                f"[dim]Populate material_cost and labor_cost fields to complete the pricing catalog.[/dim]",
                title="[bold green]debim Cost Template Generator[/bold green]",
            )
        )
    except Exception as e:
        console.print(f"[bold red]Cost Template Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_compile(args):
    manifest = resolve_default_manifest_path(args.manifest)
    output = Path(args.output)
    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    console.print(
        f"[bold green]Compiling[/bold green] {manifest} -> [cyan]{output}[/cyan]"
    )
    try:
        out_path = compile_to_ifc(manifest, output)
        file_size = out_path.stat().st_size
        console.print(
            Panel(
                f"[bold green]IFC4 Compilation Successful![/bold green]\n"
                f"[bold cyan]Output Path:[/bold cyan] {out_path}\n"
                f"[bold cyan]File Size:[/bold cyan] {file_size} bytes",
                title="[bold green]debim Compiler[/bold green]",
            )
        )
    except Exception as e:
        console.print(f"[bold red]Compilation Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_view(args):
    manifest = resolve_default_manifest_path(args.manifest)
    export_path = Path(args.export) if args.export else None

    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    try:
        watch = getattr(args, "watch", True)
        html_content = generate_viewer_html(manifest, live_reload=watch and not export_path)
        if export_path:
            export_path.parent.mkdir(parents=True, exist_ok=True)
            export_path.write_text(html_content, encoding="utf-8")
            file_size = export_path.stat().st_size
            console.print(
                Panel(
                    f"[bold green]Viewer HTML Exported Successfully![/bold green]\n"
                    f"[bold cyan]Export Path:[/bold cyan] {export_path}\n"
                    f"[bold cyan]File Size:[/bold cyan] {file_size} bytes",
                    title="[bold blue]debim 3D Viewer Export[/bold blue]",
                )
            )
            return

        url = f"http://localhost:{args.port}"
        watch_status = (
            "[bold green]Live Hot-Reload Active[/bold green] (edits in VS Code update in real-time)"
            if watch
            else "[dim]Static[/dim]"
        )
        console.print(
            Panel(
                f"[bold green]Serving 3D Web Preview at:[/bold green] [cyan bold]{url}[/cyan bold]\n"
                f"[dim]Manifest:[/dim] [yellow]{manifest}[/yellow] — {watch_status}\n"
                f"[dim]Press Ctrl+C in terminal to stop server.[/dim]",
                title="[bold blue]debim 3D Live Viewer[/bold blue]",
            )
        )
        serve_viewer(
            html_content,
            port=args.port,
            open_browser=not args.no_browser,
            manifest_path=manifest,
            watch=watch,
        )
    except Exception as e:
        console.print(f"[bold red]Viewer Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_import(args):
    ifc_path = Path(args.ifc_path)
    output = Path(args.output)
    if not ifc_path.exists():
        console.print(f"[bold red]Error:[/bold red] IFC file '{ifc_path}' not found.")
        sys.exit(1)

    console.print(f"[bold green]Importing IFC:[/bold green] {ifc_path} -> [cyan]{output}[/cyan]")
    try:
        from debim.importer import import_ifc_to_manifest
        assets_dir = output.parent / "assets" if output.parent != Path(".") else Path("assets")
        manifest = import_ifc_to_manifest(ifc_path, bake_assets=args.bake_assets, assets_dir=assets_dir)

        import json
        data = json.loads(manifest.model_dump_json(by_alias=True, exclude_none=True))
        with open(output, "w", encoding="utf-8") as f:
            yaml.safe_dump(data, f, allow_unicode=True, sort_keys=False)

        console.print(
            Panel(
                f"[bold green]IFC Successfully Imported to Declarative BIM![/bold green]\n"
                f"[cyan]Output YAML:[/cyan] {output}\n"
                f"[yellow]Elements Extracted:[/yellow] {len(manifest.elements)} elements\n"
                f"[magenta]Asset Baking Enabled:[/magenta] {args.bake_assets}\n"
                f"[dim]Run 'debim qto -m {output}' to calculate quantities & cost.[/dim]",
                title="[bold green]debim IFC Importer[/bold green]",
            )
        )
    except Exception as e:
        console.print(f"[bold red]Import Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_ingest(args):
    file_path = Path(args.file_path)
    if not file_path.exists():
        console.print(f"[bold red]Error:[/bold red] File '{file_path}' not found.")
        sys.exit(1)

    output_dir = Path(args.output_dir) if args.output_dir else None
    console.print(f"[bold green]Ingesting CAD/PDF grid & column coordinates from:[/bold green] {file_path}")

    try:
        out_path = ingest_dxf_file(file_path, output_dir=output_dir)
        console.print(
            Panel(
                f"[bold green]CAD Ingestion & Scaffolding Complete![/bold green]\n"
                f"[bold cyan]Input CAD File:[/bold cyan] {file_path}\n"
                f"[bold cyan]Emitted Manifest:[/bold cyan] {out_path}\n"
                f"[dim]Run 'debim validate {out_path}' or 'debim view {out_path}' to inspect.[/dim]",
                title="[bold green]debim CAD Ingestion Assistant[/bold green]",
            )
        )
    except Exception as e:
        console.print(f"[bold red]Ingestion Error:[/bold red]\n{e}")
        sys.exit(1)


def _resolve_target_manifest(target_str: str) -> ProjectManifest:
    resolved = resolve_default_manifest_path(target_str)
    if resolved.exists() and resolved.is_file():
        return load_manifest(resolved)

    import subprocess
    if ":" in target_str:
        rev, rel_path = target_str.split(":", 1)
    else:
        rev = target_str
        rel_path = "project.yaml"

    try:
        res = subprocess.run(
            ["git", "show", f"{rev}:{rel_path}"],
            capture_output=True,
            text=True,
            check=True,
        )
        data = yaml.safe_load(res.stdout)
        return ProjectManifest.model_validate(data)
    except Exception as e:
        raise ValueError(f"Could not load manifest from '{target_str}': {e}")


def cmd_diff(args):
    prices = Path(args.prices)
    try:
        manifest_a = _resolve_target_manifest(args.target_a)
    except Exception as e:
        console.print(f"[bold red]Error loading target A ({args.target_a}):[/bold red]\n{e}")
        sys.exit(1)

    try:
        manifest_b = _resolve_target_manifest(args.target_b)
    except Exception as e:
        console.print(f"[bold red]Error loading target B ({args.target_b}):[/bold red]\n{e}")
        sys.exit(1)

    console.print(f"[bold cyan]Comparing BIM Revisions:[/bold cyan] [yellow]{args.target_a}[/yellow] ➔ [green]{args.target_b}[/green]\n")

    qto_a = calculate_qto(manifest_a)
    qto_b = calculate_qto(manifest_b)

    cost_a = None
    cost_b = None
    if prices.exists():
        try:
            catalog = load_price_catalog(prices)
            cost_a = estimate_cost(qto_a, catalog, manifest_a)
            cost_b = estimate_cost(qto_b, catalog, manifest_b)
        except Exception:
            pass

    elems_a = {e.tag: e for e in manifest_a.elements}
    elems_b = {e.tag: e for e in manifest_b.elements}

    all_tags = sorted(list(set(elems_a.keys()) | set(elems_b.keys())))

    elem_table = Table(
        title="Element Changes Breakdown",
        show_header=True,
        header_style="bold cyan",
    )
    elem_table.add_column("Tag", style="bold yellow")
    elem_table.add_column("Class", style="white")
    elem_table.add_column("Change", justify="center")
    elem_table.add_column("Details", style="dim")

    added_count = 0
    removed_count = 0
    modified_count = 0

    for tag in all_tags:
        if tag in elems_b and tag not in elems_a:
            elem = elems_b[tag]
            elem_table.add_row(tag, elem.class_, "[bold green]+ Added[/bold green]", "New element introduced")
            added_count += 1
        elif tag in elems_a and tag not in elems_b:
            elem = elems_a[tag]
            elem_table.add_row(tag, elem.class_, "[bold red]- Removed[/bold red]", "Element deleted")
            removed_count += 1
        else:
            elem_a = elems_a[tag]
            elem_b = elems_b[tag]
            if elem_a.model_dump() != elem_b.model_dump():
                elem_table.add_row(tag, elem_b.class_, "[bold yellow]~ Modified[/bold yellow]", "Properties/geometry updated")
                modified_count += 1

    console.print(elem_table)

    def fmt_delta(val: float, unit: str = "", currency: str = "") -> str:
        prefix = f"{currency} " if currency else ""
        suffix = f" {unit}" if unit else ""
        if val > 1e-6:
            return f"[bold green]+{prefix}{val:,.3f}{suffix}[/bold green]"
        elif val < -1e-6:
            return f"[bold red]{prefix}{val:,.3f}{suffix}[/bold red]"
        else:
            return f"[dim]0.000{suffix}[/dim]"

    def fmt_cost_delta(val: float, currency: str = "THB") -> str:
        if val > 1e-2:
            return f"[bold green]+{val:,.2f} {currency}[/bold green]"
        elif val < -1e-2:
            return f"[bold red]{val:,.2f} {currency}[/bold red]"
        else:
            return f"[dim]0.00 {currency}[/dim]"

    d_conc = qto_b.total_concrete_volume - qto_a.total_concrete_volume
    d_form = qto_b.total_formwork_area - qto_a.total_formwork_area
    d_rebar = qto_b.total_rebar_weight - qto_a.total_rebar_weight

    qto_delta_table = Table(title="Quantity Take-Off (QTO) Deltas", show_header=True, header_style="bold cyan")
    qto_delta_table.add_column("Metric", style="bold yellow")
    qto_delta_table.add_column("Revision A", justify="right")
    qto_delta_table.add_column("Revision B", justify="right")
    qto_delta_table.add_column("Delta (Δ)", justify="right")

    qto_delta_table.add_row("Concrete Volume (m3)", f"{qto_a.total_concrete_volume:.3f}", f"{qto_b.total_concrete_volume:.3f}", fmt_delta(d_conc, "m3"))
    qto_delta_table.add_row("Formwork Area (m2)", f"{qto_a.total_formwork_area:.3f}", f"{qto_b.total_formwork_area:.3f}", fmt_delta(d_form, "m2"))
    qto_delta_table.add_row("Rebar Weight (kg)", f"{qto_a.total_rebar_weight:.3f}", f"{qto_b.total_rebar_weight:.3f}", fmt_delta(d_rebar, "kg"))

    console.print(qto_delta_table)

    if cost_a and cost_b:
        d_mat = cost_b.total_material_cost - cost_a.total_material_cost
        d_lab = cost_b.total_labor_cost - cost_a.total_labor_cost
        d_tot = cost_b.grand_total - cost_a.grand_total
        curr = cost_b.currency

        cost_delta_table = Table(title=f"Cost Variance Breakdown ({curr})", show_header=True, header_style="bold cyan")
        cost_delta_table.add_column("Cost Component", style="bold yellow")
        cost_delta_table.add_column("Revision A", justify="right")
        cost_delta_table.add_column("Revision B", justify="right")
        cost_delta_table.add_column("Variance (Δ)", justify="right")

        cost_delta_table.add_row("Material Cost", f"{cost_a.total_material_cost:,.2f}", f"{cost_b.total_material_cost:,.2f}", fmt_cost_delta(d_mat, curr))
        cost_delta_table.add_row("Labor Cost", f"{cost_a.total_labor_cost:,.2f}", f"{cost_b.total_labor_cost:,.2f}", fmt_cost_delta(d_lab, curr))
        cost_delta_table.add_row("Total Amount", f"{cost_a.grand_total:,.2f}", f"{cost_b.grand_total:,.2f}", fmt_cost_delta(d_tot, curr))

        console.print(cost_delta_table)

    summary_msg = (
        f"[bold yellow]Elements Changed:[/bold yellow] [green]+{added_count} Added[/green] | "
        f"[red]-{removed_count} Removed[/red] | [yellow]~{modified_count} Modified[/yellow]\n"
        f"[bold cyan]Net QTO Deltas:[/bold cyan] Δ Concrete: {d_conc:+.3f} m3, Δ Formwork: {d_form:+.3f} m2, Δ Rebar: {d_rebar:+.3f} kg"
    )
    if cost_a and cost_b:
        summary_msg += f"\n[bold green]Net Budget Variance:[/bold green] {cost_b.grand_total - cost_a.grand_total:+,.2f} {cost_b.currency}"

    console.print(Panel(summary_msg, title="[bold green]Diff Summary[/bold green]"))


def cmd_scaffold_element(args):
    class_name = args.class_name
    if not class_name.startswith("Ifc"):
        class_name = f"Ifc{class_name[0].upper()}{class_name[1:]}"

    dry_run = getattr(args, "dry_run", True)
    output_dir = Path(args.output_dir) if args.output_dir else None

    console.print(
        Panel(
            f"[bold cyan]debim Element Scaffolder[/bold cyan]\n"
            f"[bold yellow]Element Class:[/bold yellow] {class_name}\n"
            f"[bold yellow]Placement Type:[/bold yellow] {args.placement_type}\n"
            f"[bold yellow]Default Layer:[/bold yellow] {args.layer or 'derived'}\n"
            f"[bold yellow]Dry Run Mode:[/bold yellow] {dry_run}",
            title="[bold green]scaffold element[/bold green]",
        )
    )

    artifacts = scaffold_element(
        class_name=class_name,
        placement_type=args.placement_type,
        layer=args.layer,
        dry_run=dry_run,
        output_dir=output_dir,
    )

    if dry_run:
        console.print("[bold magenta]=== 1. Generated Pydantic v2 Model (for schema.py) ===[/bold magenta]")
        console.print(artifacts["schema_code"])

        console.print("\n[bold magenta]=== 2. Generated Resolver Logic (for resolver.py) ===[/bold magenta]")
        console.print(artifacts["resolver_code"])

        console.print("\n[bold magenta]=== 3. Generated QTO Formula (for qto.py) ===[/bold magenta]")
        console.print(artifacts["qto_code"])

        console.print("\n[bold magenta]=== 4. Generated Pytest Skeleton ===[/bold magenta]")
        console.print(artifacts["test_code"])

        console.print(
            Panel(
                "[bold green]Dry-run complete![/bold green] Pass [cyan]--no-dry-run -o <output_dir>[/cyan] to write files to disk.",
                style="green",
            )
        )
    else:
        if output_dir:
            console.print(
                Panel(
                    f"[bold green]Scaffold artifacts successfully written to:[/bold green] [cyan]{output_dir}[/cyan]",
                    style="green",
                )
            )
        else:
            console.print("[bold red]Error:[/bold red] --no-dry-run requires --output-dir (-o) specified.")
            sys.exit(1)


def cmd_draw_view(args):
    manifest = resolve_default_manifest_path(args.manifest)
    output = Path(args.output)
    sheets = Path(args.sheets) if args.sheets else None

    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    console.print(
        f"[bold green]Generating 2D Blueprint HTML Viewer:[/bold green] {manifest} -> [cyan]{output}[/cyan]"
    )
    try:
        from debim.draw import export_2d_viewer

        out_path = export_2d_viewer(
            manifest_or_resolved=manifest,
            output_path=output,
            sheets_dir_or_configs=sheets,
        )

        file_size = out_path.stat().st_size
        console.print(
            Panel(
                f"[bold green]2D Blueprint Viewer Export Successful![/bold green]\n"
                f"[bold cyan]Output HTML File:[/bold cyan] {out_path}\n"
                f"[bold cyan]File Size:[/bold cyan] {file_size:,} bytes\n"
                f"[dim]Open in web browser for interactive 2D sheet inspection, pan-zoom & ruler measuring.[/dim]",
                title="[bold green]debim 2D Viewer Engine[/bold green]",
            )
        )
    except Exception as e:
        console.print(f"[bold red]2D Viewer Export Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_draw_export_dxf(args):
    manifest = resolve_default_manifest_path(args.manifest)
    output = Path(args.output)

    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    console.print(
        f"[bold green]Exporting 2D CAD DXF:[/bold green] {manifest} -> [cyan]{output}[/cyan]"
    )
    try:
        from debim.draw import export_2d_dxf

        manifest_obj = load_manifest(manifest)
        out_path = export_2d_dxf(
            manifest_or_resolved=manifest_obj,
            output_path=output,
            storey_id=args.storey,
            cut_offset_z=args.cut_offset_z,
            units=args.units,
        )

        import ezdxf
        dxf_doc = ezdxf.readfile(out_path)
        entity_count = len(list(dxf_doc.modelspace()))
        layer_names = [layer.dxf.name for layer in dxf_doc.layers]

        console.print(
            Panel(
                f"[bold green]2D AutoCAD DXF Export Successful![/bold green]\n"
                f"[bold cyan]Output File:[/bold cyan] {out_path}\n"
                f"[bold cyan]Entities Exported:[/bold cyan] {entity_count}\n"
                f"[bold cyan]CAD Units:[/bold cyan] {args.units.upper()} ($INSUNITS={dxf_doc.header.get('$INSUNITS', 4)})\n"
                f"[bold cyan]Exported Layers:[/bold cyan] {', '.join(layer_names)}",
                title="[bold green]debim 2D DXF Exporter Engine[/bold green]",
            )
        )
    except Exception as e:
        console.print(f"[bold red]DXF Export Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_draw_generate_vqa(args):
    manifest = Path(args.manifest) if args.manifest else None
    output = Path(args.output)
    image_dir = Path(args.image_dir)

    console.print(f"[bold green]Generating VQA MCQ Dataset:[/bold green] -> [cyan]{output}[/cyan]")
    try:
        if str(Path.cwd()) not in sys.path:
            sys.path.insert(0, str(Path.cwd()))
        from tools.generate_vqa_dataset import generate_vqa_dataset

        ds_data = generate_vqa_dataset(
            manifest_path=manifest,
            output_dataset_path=output,
            output_image_dir=image_dir,
            seed=args.seed,
        )

        q_count = ds_data.get("question_count", 0)
        cat_counts = ds_data.get("category_counts", {})
        cat_summary = ", ".join([f"{k}: {v}" for k, v in cat_counts.items()])

        console.print(
            Panel(
                f"[bold green]VQA Dataset Generation Successful![/bold green]\n"
                f"[bold cyan]Dataset Path:[/bold cyan] {output}\n"
                f"[bold cyan]Total Questions:[/bold cyan] {q_count}\n"
                f"[bold cyan]Category Breakdown:[/bold cyan] {cat_summary}\n"
                f"[bold cyan]Image Asset Path:[/bold cyan] {ds_data.get('image_path', '')}",
                title="[bold green]debim VQA Dataset Generator[/bold green]",
            )
        )
    except Exception as e:
        console.print(f"[bold red]VQA Dataset Generation Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_draw_benchmark_vqa(args):
    dataset = Path(args.dataset)
    output = Path(args.output)
    markdown_output = Path(args.markdown_output)

    console.print(f"[bold green]Running VQA Cognitive Benchmark:[/bold green] [yellow]{dataset}[/yellow]")
    try:
        if str(Path.cwd()) not in sys.path:
            sys.path.insert(0, str(Path.cwd()))
        from tools.benchmark_vqa_gemini import run_vqa_benchmark

        report = run_vqa_benchmark(
            dataset_path=dataset,
            api_key=args.api_key,
            model_name=args.model,
            mock_eval=args.mock_eval,
            rpm=args.rpm,
            tpm=args.tpm,
            thinking_budget=args.thinking_budget,
            output_report_path=output,
            output_markdown_path=markdown_output,
        )

        cat_str = ", ".join([f"{k}: {v:.1f}%" for k, v in report.category_accuracy.items()])

        console.print(
            Panel(
                f"[bold green]VQA Cognitive Benchmark Complete![/bold green]\n"
                f"[bold cyan]AI Model:[/bold cyan] {report.model_name}\n"
                f"[bold cyan]Overall Accuracy:[/bold cyan] [bold green]{report.overall_accuracy_pct:.2f}%[/bold green] ({report.correct_count}/{report.total_questions})\n"
                f"[bold cyan]Category Accuracy:[/bold cyan] {cat_str}\n"
                f"[bold cyan]Diagnostic Failure Cases:[/bold cyan] {len(report.diagnostic_failures)}\n"
                f"[bold cyan]JSON Report:[/bold cyan] {output}\n"
                f"[bold cyan]Markdown Report:[/bold cyan] {markdown_output}",
                title="[bold green]debim VQA Cognitive Benchmark Engine[/bold green]",
            )
        )
    except Exception as e:
        console.print(f"[bold red]VQA Benchmark Execution Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_spec_build(args):
    manifest = resolve_default_manifest_path(args.manifest)
    specs = Path(args.specs) if args.specs else None
    output = Path(args.output)

    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    console.print(f"[bold blue]Compiling Specification Book for:[/bold blue] {manifest}")
    if specs:
        console.print(f"[bold blue]Using Specs Path:[/bold blue] {specs}")

    try:
        from debim.spec.builder import build_specification_book

        book = build_specification_book(
            project_manifest=manifest,
            spec_manifest=specs,
            title=args.title,
        )

        saved_path = book.save(output_path=output, format=args.format)
        file_size = saved_path.stat().st_size

        table = Table(
            title="Specification Book Summary",
            show_header=True,
            header_style="bold cyan",
        )
        table.add_column("MasterFormat Division", style="bold yellow")
        table.add_column("Division Title", style="white")
        table.add_column("Specs Count", justify="right", style="green")

        for div_code, div_data in book.divisions.items():
            table.add_row(
                f"Division {div_code}",
                div_data["title"],
                str(len(div_data["specs"])),
            )

        console.print(table)

        summary_text = (
            f"[bold cyan]Project Name:[/bold cyan] {book.project_name} ({book.project_id})\n"
            f"[bold cyan]Active Specifications Included:[/bold cyan] {len(book.included_specs)} (zero unreferenced bloat)\n"
            f"[bold cyan]MasterFormat Divisions Covered:[/bold cyan] {len(book.divisions)}\n"
            f"[bold green]Output Specification Book:[/bold green] [cyan]{saved_path}[/cyan] ({file_size:,} bytes)"
        )
        console.print(
            Panel(
                summary_text,
                title="[bold green]debim Spec Book Build Complete[/bold green]",
                style="green",
            )
        )

    except SystemExit:
        raise
    except Exception as e:
        console.print(f"[bold red]Specification Book Build Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_spec_to_word(args):
    input_file = Path(args.input_file)
    if not input_file.exists():
        console.print(f"[bold red]Error:[/bold red] Input file '{input_file}' not found.")
        sys.exit(1)

    out_p = Path(args.output) if args.output else input_file.with_suffix(".docx")
    console.print(f"[bold blue]Converting specification book to MS Word:[/bold blue] {input_file} -> {out_p}")

    try:
        from debim.spec.docx_exporter import markdown_to_docx

        md_content = input_file.read_text(encoding="utf-8")
        saved_path = markdown_to_docx(md_content, out_p, title=args.title, font_name=args.font)
        file_size = saved_path.stat().st_size

        console.print(
            Panel(
                f"[bold green]MS Word Document Exported Successfully![/bold green]\n"
                f"[bold cyan]Input File:[/bold cyan] {input_file}\n"
                f"[bold cyan]Output File:[/bold cyan] [cyan]{saved_path}[/cyan] ({file_size:,} bytes)\n"
                f"[bold cyan]Font Family:[/bold cyan] {args.font}",
                title="[bold green]debim Spec Word Export[/bold green]",
                style="green",
            )
        )
    except Exception as e:
        console.print(f"[bold red]Word Conversion Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_spec_audit(args):
    manifest = resolve_default_manifest_path(args.manifest)
    specs = Path(args.specs) if args.specs else None

    if not manifest.exists():
        console.print(f"[bold red]Error:[/bold red] Manifest '{manifest}' not found.")
        sys.exit(1)

    console.print(f"[bold blue]Auditing Specifications for:[/bold blue] {manifest}")
    if specs:
        console.print(f"[bold blue]Using Specs Path:[/bold blue] {specs}")

    try:
        from debim.spec import audit_project_specs, load_spec_manifest

        specs_manifest_obj = load_spec_manifest(specs) if specs else None
        res = audit_project_specs(manifest, specs_manifest_obj)

        table = Table(
            title="Material Specification Discrepancy Audit",
            show_header=True,
            header_style="bold cyan",
        )
        table.add_column("Material / Spec ID", style="bold white")
        table.add_column("Status", justify="center")
        table.add_column("Referenced Element Tags", style="dim")

        for mat_id in res.missing_specs:
            tags = ", ".join(res.element_material_map.get(mat_id, [])) or "(manifest declared)"
            table.add_row(
                mat_id,
                "[bold red]MISSING SPEC[/bold red]",
                tags,
            )

        for spec_id in res.unused_specs:
            table.add_row(
                spec_id,
                "[bold yellow]UNUSED / REDUNDANT[/bold yellow]",
                "(none)",
            )

        compliant_specs = set(res.referenced_materials).intersection(set(res.spec_ids))
        for mat_id in sorted(compliant_specs):
            tags = ", ".join(res.element_material_map.get(mat_id, [])) or "(manifest declared)"
            table.add_row(
                mat_id,
                "[bold green]MATCHED / OK[/bold green]",
                tags,
            )

        console.print(table)

        summary_text = (
            f"[bold]Referenced Materials:[/bold] {len(res.referenced_materials)}\n"
            f"[bold]Specification Packages:[/bold] {len(res.spec_ids)}\n"
            f"[bold red]Missing Specs:[/bold red] {len(res.missing_specs)}\n"
            f"[bold yellow]Unused / Redundant Specs:[/bold yellow] {len(res.unused_specs)}\n"
            f"[bold green]Compliant Specs:[/bold green] {len(compliant_specs)}"
        )

        title_style = "bold green" if res.is_compliant else "bold red"
        status_label = "COMPLIANT PASS" if res.is_compliant else "DISCREPANCIES DETECTED"
        console.print(
            Panel(
                summary_text,
                title=f"[{title_style}]Audit Summary: {status_label}[/{title_style}]",
            )
        )

        if args.strict and not res.is_compliant:
            sys.exit(1)

    except SystemExit:
        raise
    except Exception as e:
        console.print(f"[bold red]Specification Audit Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_mcp(args):
    from debim.mcp import run_mcp_server
    run_mcp_server()


def cmd_spec_add(args):
    specs_dir = Path(args.specs_dir)
    cache_dir = Path(args.cache_dir) if args.cache_dir else None

    console.print(f"[bold cyan]Fetching specification package:[/bold cyan] [yellow]{args.pkg_name}[/yellow]...")
    try:
        client = SpecRegistryClient(cache_dir=cache_dir)
        installed_path = client.install_package(args.pkg_name, specs_dir=specs_dir, force=args.force)
        pkg = client.fetch_package(args.pkg_name, force=False)

        console.print(
            Panel(
                f"[bold green]Specification Package Installed Successfully![/bold green]\n"
                f"[bold cyan]Package Name:[/bold cyan] {pkg.name}\n"
                f"[bold cyan]Version:[/bold cyan] {pkg.version}\n"
                f"[bold cyan]Manufacturer:[/bold cyan] {pkg.manufacturer or 'N/A'}\n"
                f"[bold cyan]Installed Path:[/bold cyan] {installed_path}",
                title="[bold green]debim Spec Package Manager[/bold green]",
            )
        )
    except Exception as e:
        console.print(f"[bold red]Error installing specification package '{args.pkg_name}':[/bold red]\n{e}")
        sys.exit(1)


def cmd_spec_list(args):
    specs_dir = Path(args.specs_dir)
    client = SpecRegistryClient()
    installed = client.list_installed_packages(specs_dir=specs_dir)

    if not installed:
        console.print(
            f"[yellow]No specification packages installed in '[cyan]{specs_dir}[/cyan]'.[/yellow]"
        )
        console.print("[dim]Use 'debim spec add <pkg_name>' to install material specifications.[/dim]")
        return

    table = Table(
        title=f"Installed Material Specification Packages ({specs_dir})",
        show_header=True,
        header_style="bold cyan",
    )
    table.add_column("Package Name", style="bold yellow")
    table.add_column("Version", justify="center", style="green")
    table.add_column("Category", style="magenta")
    table.add_column("Manufacturer", style="white")
    table.add_column("Path", style="dim")

    for item in installed:
        table.add_row(
            item["name"],
            item["version"],
            item["category"],
            item["manufacturer"],
            item["path"],
        )

    console.print(table)


def cmd_docs_directory(args):
    output = Path(args.output)
    console.print(f"[bold green]Generating BIM Class Directory Webpage:[/bold green] -> [cyan]{output}[/cyan]")
    try:
        from debim.docs import export_class_directory

        out_path = export_class_directory(output)
        file_size = out_path.stat().st_size
        console.print(
            Panel(
                f"[bold green]Class Directory Web Application Exported Successfully![/bold green]\n"
                f"[bold cyan]Output HTML File:[/bold cyan] {out_path}\n"
                f"[bold cyan]File Size:[/bold cyan] {file_size:,} bytes\n"
                f"[dim]Open in web browser for interactive 6D BIM class inspection & encyclopedia search.[/dim]",
                title="[bold green]debim Documentation Engine[/bold green]",
            )
        )
    except Exception as e:
        console.print(f"[bold red]Docs Directory Export Error:[/bold red]\n{e}")
        sys.exit(1)


def _generate_docs_bundle(output_dir: Path, schema_version: str = "IFC4") -> Path:
    import json
    from debim.schema_registry import SchemaEntityRegistry

    output_dir.mkdir(parents=True, exist_ok=True)
    registry = SchemaEntityRegistry(schema_version)

    entities_dict = {}
    decl_names = sorted([decl.name() for decl in registry.schema.entities()])

    for name in decl_names:
        info = registry.get_entity_info(name)
        entities_dict[name] = {
            "name": info.name,
            "schema_version": info.schema_version,
            "is_abstract": info.is_abstract,
            "supertype": info.supertype,
            "subtypes": info.subtypes,
            "inheritance_chain": info.inheritance_chain,
            "is_product": info.is_product,
            "is_element": info.is_element,
            "is_spatial": info.is_spatial,
            "required_attributes": [a.model_dump() for a in info.required_attributes],
            "attributes": [a.model_dump() for a in info.attributes],
            "allowed_predefined_types": info.allowed_predefined_types,
            "default_psets": info.default_psets,
        }

    json_path = output_dir / "entities.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(entities_dict, f, indent=2, ensure_ascii=False)

    from debim.docs.generator import export_class_directory

    index_path = output_dir / "index.html"
    export_class_directory(index_path)

    farnsworth_src = Path("examples/farnsworth_house/viewer.html")
    if farnsworth_src.exists():
        import shutil
        try:
            shutil.copyfile(farnsworth_src, output_dir / "farnsworth.html")
            ex_dir = output_dir / "examples" / "farnsworth_house"
            ex_dir.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(farnsworth_src, ex_dir / "viewer.html")
        except Exception:
            pass

    content = index_path.read_text(encoding="utf-8")
    if "debim Interactive Class Directory" not in content:
        content = content.replace("<title>debim BIM Class Directory & Encyclopedia</title>", f"<title>debim Interactive Class Directory ({schema_version})</title>")
    if schema_version not in content:
        content = content.replace('<span class="badge-schema" id="badge-schema">IFC4</span>', f'<span class="badge-schema" id="badge-schema">{schema_version}</span>')
    index_path.write_text(content, encoding="utf-8")

    return output_dir


def cmd_docs_build(args):
    output_dir = Path(args.output_dir)
    console.print(f"[bold cyan]Building debim Class Directory ({args.schema_version}):[/bold cyan] -> [yellow]{output_dir}[/yellow]")
    try:
        out_p = _generate_docs_bundle(output_dir=output_dir, schema_version=args.schema_version)
        files = list(out_p.glob("*"))
        index_size = (out_p / "index.html").stat().st_size
        json_size = (out_p / "entities.json").stat().st_size

        console.print(
            Panel(
                f"[bold green]Interactive Class Directory Generated Successfully![/bold green]\n"
                f"[bold cyan]Bundle Path:[/bold cyan] {out_p.resolve()}\n"
                f"[bold cyan]Static Files Emitted:[/bold cyan] {len(files)} files\n"
                f"[bold cyan]index.html Size:[/bold cyan] {index_size:,} bytes\n"
                f"[bold cyan]entities.json Size:[/bold cyan] {json_size:,} bytes\n"
                f"[dim]Run 'debim docs serve' to launch preview HTTP server.[/dim]",
                title="[bold green]debim Documentation Engine[/bold green]",
            )
        )
    except Exception as e:
        console.print(f"[bold red]Docs Build Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_docs_serve(args):
    import http.server
    import socketserver
    import webbrowser

    docs_dir = Path(args.docs_dir)
    if not docs_dir.exists() or not (docs_dir / "index.html").exists():
        console.print(f"[bold yellow]Documentation bundle not found in '{docs_dir}'. Building now...[/bold yellow]")
        _generate_docs_bundle(output_dir=docs_dir)

    url = f"http://localhost:{args.port}"
    console.print(
        Panel(
            f"[bold green]Serving debim Class Directory at:[/bold green] [cyan bold]{url}[/cyan bold]\n"
            f"[dim]Docs Directory:[/dim] [yellow]{docs_dir.resolve()}[/yellow]\n"
            f"[dim]Press Ctrl+C in terminal to stop server.[/dim]",
            title="[bold blue]debim Docs Live Preview[/bold blue]",
        )
    )

    class Handler(http.server.SimpleHTTPRequestHandler):
        def __init__(self, *a, **kw):
            super().__init__(*a, directory=str(docs_dir.resolve()), **kw)

        def log_message(self, format, *a):
            pass

    if not args.no_browser:
        webbrowser.open(url)

    try:
        with socketserver.TCPServer(("0.0.0.0", args.port), Handler) as httpd:
            httpd.serve_forever()
    except KeyboardInterrupt:
        console.print("\n[dim]Stopped docs preview server.[/dim]")


def cmd_docs_exemplar(args):
    try:
        output_path = Path(args.output) if args.output else None
        manifest_obj = generate_exemplar_manifest(class_name=args.class_name)
        dict_data = manifest_obj.model_dump(by_alias=True, exclude_none=True)

        if args.format.lower() == "json":
            import json
            out_str = json.dumps(dict_data, indent=2, ensure_ascii=False) + "\n"
        else:
            out_str = yaml.safe_dump(dict_data, allow_unicode=True, sort_keys=False)

        if output_path:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(out_str, encoding="utf-8")
            console.print(
                Panel(
                    f"[bold green]Exemplar Manifest Generated Successfully![/bold green]\n"
                    f"[bold cyan]Target Class:[/bold cyan] {args.class_name or 'ALL'}\n"
                    f"[bold cyan]Elements Generated:[/bold cyan] {len(manifest_obj.elements)}\n"
                    f"[bold cyan]Saved Path:[/bold cyan] {output_path}",
                    title="[bold green]debim Exemplar Generator[/bold green]",
                )
            )
        else:
            sys.stdout.write(out_str)

    except Exception as e:
        console.print(f"[bold red]Exemplar Generation Error:[/bold red]\n{e}")
        sys.exit(1)


def cmd_docs_export_data(args):
    output_dir = Path(args.output_dir)
    try:
        output_dir.mkdir(parents=True, exist_ok=True)

        full_manifest = generate_exemplar_manifest(class_name="ALL")
        full_yaml_path = output_dir / "exemplar_all.debim.yaml"
        full_yaml_path.write_text(
            yaml.safe_dump(full_manifest.model_dump(by_alias=True, exclude_none=True), allow_unicode=True, sort_keys=False),
            encoding="utf-8",
        )

        exemplars_dir = output_dir / "exemplars"
        exemplars_dir.mkdir(exist_ok=True)
        supported = get_supported_exemplar_classes()
        for cls_name in supported:
            elem_data = generate_exemplar_element(cls_name)
            cls_file = exemplars_dir / f"{cls_name.lower()}.yaml"
            cls_file.write_text(yaml.safe_dump(elem_data, allow_unicode=True, sort_keys=False), encoding="utf-8")

        specs_dir = output_dir / "specs_thai"
        specs_dir.mkdir(exist_ok=True)
        for code, clause in THAI_SPEC_LIBRARY.items():
            code_slug = code.replace(" ", "_")
            spec_file = specs_dir / f"masterformat_{code_slug}.md"
            spec_file.write_text(clause.to_markdown(), encoding="utf-8")

        console.print(
            Panel(
                f"[bold green]Exemplar & Thai Specs Data Exported Successfully![/bold green]\n"
                f"[bold cyan]Output Directory:[/bold cyan] {output_dir}\n"
                f"[bold cyan]Classes Exported:[/bold cyan] {len(supported)} individual YAML exemplars\n"
                f"[bold cyan]Thai Specs Exported:[/bold cyan] {len(THAI_SPEC_LIBRARY)} MasterFormat division Markdown specs",
                title="[bold green]debim Docs Exporter[/bold green]",
            )
        )

    except Exception as e:
        console.print(f"[bold red]Docs Export Error:[/bold red]\n{e}")
        sys.exit(1)


# --- Argparse Parser Construction ---

def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="debim",
        description="Minimal Declarative BIM (Building-as-Code) engine",
    )
    parser.add_argument("-v", "--version", action="store_true", help="Show debim version")
    parser.add_argument("-q", "--quiet", action="store_true", help="Suppress update check and non-essential output")

    subparsers = parser.add_subparsers(dest="subcommand", help="Available subcommands")

    # init
    p_init = subparsers.add_parser("init", help="Create a new project scaffold")
    p_init.add_argument("name", help="Name of the new project directory")
    p_init.set_defaults(func=cmd_init)

    # validate
    p_validate = subparsers.add_parser("validate", help="Validate project.yaml schema, topology, and bSDD Psets")
    p_validate.add_argument("positional_manifest", nargs="?", help="Positional path to project manifest")
    p_validate.add_argument("-m", "--manifest", help="Path to project manifest")
    bsdd_group = p_validate.add_mutually_exclusive_group()
    bsdd_group.add_argument("--bsdd", dest="bsdd", action="store_true", default=True, help="Perform buildingSMART bSDD Property Set validation")
    bsdd_group.add_argument("--no-bsdd", dest="bsdd", action="store_false", help="Disable bSDD validation")
    p_validate.add_argument("--strict-psets", action="store_true", default=False, help="Fail validation if non-compliant Pset types or values are found")
    p_validate.set_defaults(func=cmd_validate)

    # schema
    p_schema = subparsers.add_parser("schema", help="Export debim language specification as JSON Schema")
    p_schema.add_argument("-o", "--output", help="Target file path to write JSON schema")
    p_schema.add_argument("--indent", type=int, default=2, help="JSON indentation level")
    p_schema.add_argument("--vscode", action="store_true", help="Configure .vscode/settings.json with schema mapping")
    p_schema.set_defaults(func=cmd_schema)

    # split
    p_split = subparsers.add_parser("split", help="Split monolithic project.yaml into modular files")
    p_split.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_split.add_argument("-o", "--output-dir", help="Target output directory")
    p_split.add_argument("-b", "--by", default="system", help="Decomposition strategy: 'system' or 'storey'")
    backup_group = p_split.add_mutually_exclusive_group()
    backup_group.add_argument("--backup", dest="backup", action="store_true", default=True, help="Create backup of original manifest")
    backup_group.add_argument("--no-backup", dest="backup", action="store_false", help="Do not create backup")
    p_split.add_argument("--dry-run", action="store_true", help="Preview decomposition without writing files")
    p_split.set_defaults(func=cmd_split)

    # bundle
    p_bundle = subparsers.add_parser("bundle", help="Bundle modular project manifest into single file")
    p_bundle.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_bundle.add_argument("-o", "--output", help="Output path for bundled single-file manifest")
    p_bundle.set_defaults(func=cmd_bundle)

    # test
    p_test = subparsers.add_parser("test", help="Run building law & compliance test suite")
    p_test.add_argument("-t", "--tests", dest="test_dir", default="tests", help="Directory containing compliance tests")
    p_test.add_argument("-v", "--verbose", action="store_true", help="Display verbose pytest output")
    p_test.set_defaults(func=cmd_test)

    # summary / info
    p_summary = subparsers.add_parser("summary", help="Print high-level project summary table")
    p_summary.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_summary.add_argument("-p", "--prices", default="prices.json", help="Path to price catalog")
    p_summary.set_defaults(func=cmd_summary)

    p_info = subparsers.add_parser("info", help=argparse.SUPPRESS)
    p_info.add_argument("-m", "--manifest", default="project.yaml")
    p_info.add_argument("-p", "--prices", default="prices.json")
    p_info.set_defaults(func=cmd_summary)

    # qto
    p_qto = subparsers.add_parser("qto", help="Calculate Quantitative Take-Off (QTO)")
    p_qto.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_qto.set_defaults(func=cmd_qto)

    # cost
    p_cost = subparsers.add_parser("cost", help="Calculate cost estimate or generate price catalog template")
    p_cost.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_cost.add_argument("-p", "--prices", default="prices.json", help="Path or '-' to read price catalog")
    p_cost.add_argument("-o", "--output", help="Path to export BOQ CSV")

    cost_subparsers = p_cost.add_subparsers(dest="cost_subcommand")
    p_cost_template = cost_subparsers.add_parser("template", help="Generate price catalog template")
    p_cost_template.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_cost_template.add_argument("-o", "--output", help="Path for generated template file")
    p_cost_template.add_argument("--format", default="yaml", help="Output format (yaml or json)")
    mod_group = p_cost_template.add_mutually_exclusive_group()
    mod_group.add_argument("--modular", dest="modular", action="store_true", help="Generate modular folder structure")
    mod_group.add_argument("--no-modular", dest="modular", action="store_false", default=False)
    p_cost_template.set_defaults(func=cmd_cost_template)
    p_cost.set_defaults(func=cmd_cost)

    # boq
    p_boq = subparsers.add_parser("boq", help="Generate Bill of Quantities (BOQ) with Thai Public Works pricing rules")
    p_boq.add_argument("positional_manifest", nargs="?", help="Positional path to project manifest")
    p_boq.add_argument("-m", "--manifest", help="Path to project manifest")
    p_boq.add_argument("--rules", default="cg-v480", help="Pricing rules standard (e.g. 'cg-v480' for Thai CGD v480)")
    p_boq.add_argument("-p", "--prices", help="Path to price catalog file")
    p_boq.add_argument("--province", help="Select provincial price index multiplier")
    p_boq.add_argument("--export-excel", help="Target Excel (.xlsx) file path to export BOQ sheets")
    p_boq.set_defaults(func=cmd_boq)

    # compile
    p_compile = subparsers.add_parser("compile", help="Compile project.yaml to IFC4 format")
    p_compile.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_compile.add_argument("-o", "--output", default="dist/model.ifc", help="Output IFC file path")
    p_compile.set_defaults(func=cmd_compile)

    # view
    p_view = subparsers.add_parser("view", help="Launch local 3D preview server or export HTML")
    p_view.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_view.add_argument("-p", "--port", type=int, default=8000, help="Port to serve 3D viewer")
    p_view.add_argument("--no-browser", action="store_true", help="Do not open web browser automatically")
    watch_group = p_view.add_mutually_exclusive_group()
    watch_group.add_argument("-w", "--watch", dest="watch", action="store_true", default=True, help="Enable automatic live hot-reload")
    watch_group.add_argument("--no-watch", dest="watch", action="store_false", help="Disable live hot-reload")
    p_view.add_argument("-e", "--export", help="Export standalone 3D/2D HTML viewer file")
    p_view.set_defaults(func=cmd_view)

    # import
    p_import = subparsers.add_parser("import", help="Import an IFC file and convert to project.yaml")
    p_import.add_argument("ifc_path", help="Path to input IFC file (.ifc)")
    p_import.add_argument("-o", "--output", default="project.yaml", help="Output project manifest path")
    bake_group = p_import.add_mutually_exclusive_group()
    bake_group.add_argument("--bake-assets", dest="bake_assets", action="store_true", help="Automated GLB asset baker")
    bake_group.add_argument("--no-bake-assets", dest="bake_assets", action="store_false", default=False)
    p_import.set_defaults(func=cmd_import)

    # ingest
    p_ingest = subparsers.add_parser("ingest", help="Extract grid labels and dimensions from DXF/CAD into project.debim")
    p_ingest.add_argument("file_path", help="Path to input DXF/CAD drawing file")
    p_ingest.add_argument("-o", "--out", dest="output_dir", help="Output directory path for generated project.debim")
    p_ingest.set_defaults(func=cmd_ingest)

    # diff
    p_diff = subparsers.add_parser("diff", help="Git-aware manifest comparison engine")
    p_diff.add_argument("target_a", help="First target (file path or git rev)")
    p_diff.add_argument("target_b", help="Second target (file path or git rev)")
    p_diff.add_argument("-p", "--prices", default="prices.json", help="Path to price catalog")
    p_diff.set_defaults(func=cmd_diff)

    # scaffold
    p_scaffold = subparsers.add_parser("scaffold", help="Developer scaffolding commands")
    scaffold_subparsers = p_scaffold.add_subparsers(dest="scaffold_subcommand")
    p_scaffold_elem = scaffold_subparsers.add_parser("element", help="Generate code skeleton for a new element class")
    p_scaffold_elem.add_argument("class_name", help="BIM Element class name")
    p_scaffold_elem.add_argument("-p", "--placement-type", default="grid", help="Placement geometry type")
    p_scaffold_elem.add_argument("-l", "--layer", help="Default hierarchical layer path")
    dry_group = p_scaffold_elem.add_mutually_exclusive_group()
    dry_group.add_argument("--dry-run", dest="dry_run", action="store_true", default=True, help="Print generated code to stdout")
    dry_group.add_argument("--no-dry-run", dest="dry_run", action="store_false", help="Write generated code to disk")
    p_scaffold_elem.add_argument("-o", "--output-dir", help="Directory path to emit generated code")
    p_scaffold_elem.set_defaults(func=cmd_scaffold_element)

    # draw
    p_draw = subparsers.add_parser("draw", help="2D Architectural Blueprint & CAD export commands")
    draw_subparsers = p_draw.add_subparsers(dest="draw_subcommand")

    p_draw_view = draw_subparsers.add_parser("view", help="Generate standalone interactive 2D HTML blueprint viewer")
    p_draw_view.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_draw_view.add_argument("--sheets", help="Directory containing sheet configs or single sheet file")
    p_draw_view.add_argument("-o", "--output", default="dist/viewer_2d.html", help="Output 2D HTML viewer file path")
    p_draw_view.set_defaults(func=cmd_draw_view)

    p_draw_dxf = draw_subparsers.add_parser("export-dxf", help="Export 2D floor plan to AutoCAD DXF")
    p_draw_dxf.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_draw_dxf.add_argument("-o", "--output", default="dist/plan.dxf", help="Output DXF file path")
    p_draw_dxf.add_argument("-s", "--storey", help="Target storey ID to slice")
    p_draw_dxf.add_argument("-u", "--units", default="mm", help="Output DXF length units ('mm' or 'm')")
    p_draw_dxf.add_argument("--cut-offset-z", type=float, default=1.20, help="Cut-plane elevation offset above storey")
    p_draw_dxf.set_defaults(func=cmd_draw_export_dxf)

    p_draw_vqa = draw_subparsers.add_parser("generate-vqa", help="Generate 2D Construction Blueprint VQA MCQ Dataset")
    p_draw_vqa.add_argument("-m", "--manifest", help="Path to project manifest")
    p_draw_vqa.add_argument("-o", "--output", default="dist/vqa_dataset.json", help="Output VQA dataset JSON path")
    p_draw_vqa.add_argument("--image-dir", default="dist/vqa_images", help="Directory to export rendered 2D SVG images")
    p_draw_vqa.add_argument("--seed", type=int, default=42, help="Random seed")
    p_draw_vqa.set_defaults(func=cmd_draw_generate_vqa)

    p_draw_bench = draw_subparsers.add_parser("benchmark-vqa", help="Run 2D Construction Blueprint VQA Cognitive Benchmark")
    p_draw_bench.add_argument("-d", "--dataset", default="dist/vqa_dataset.json", help="Path to VQA dataset JSON")
    p_draw_bench.add_argument("-k", "--api-key", help="Gemini API key")
    p_draw_bench.add_argument("-m", "--model", default="models/gemini-3.8-live-extended-thinking", help="Gemini multimodal model")
    p_draw_bench.add_argument("--thinking-budget", type=int, default=2048, help="Thinking budget")
    p_draw_bench.add_argument("--mock-eval", action="store_true", help="Perform mock evaluation")
    p_draw_bench.add_argument("-o", "--output", default="dist/benchmark_vqa_report.json", help="Output JSON report path")
    p_draw_bench.add_argument("--markdown-output", default="dist/benchmark_vqa_report.md", help="Output Markdown report path")
    p_draw_bench.add_argument("--rpm", type=int, default=60, help="Requests per minute rate limit")
    p_draw_bench.add_argument("--tpm", type=int, default=65000, help="Tokens per minute budget limit")
    p_draw_bench.set_defaults(func=cmd_draw_benchmark_vqa)

    # spec
    p_spec = subparsers.add_parser("spec", help="Material specification management & auditing")
    spec_subparsers = p_spec.add_subparsers(dest="spec_subcommand")

    p_spec_build = spec_subparsers.add_parser("build", help="Compile a Specification Book")
    p_spec_build.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_spec_build.add_argument("-s", "--specs", help="Path to material specifications manifest or directory")
    p_spec_build.add_argument("-o", "--output", default="dist/specifications.md", help="Output file path")
    p_spec_build.add_argument("-f", "--format", help="Output format override")
    p_spec_build.add_argument("-t", "--title", default="ARCHITECTURAL MATERIAL SPECIFICATIONS BOOK", help="Title of Specification Book")
    p_spec_build.set_defaults(func=cmd_spec_build)

    p_spec_word = spec_subparsers.add_parser("to-word", help="Convert Markdown specification book to MS Word")
    p_spec_word.add_argument("input_file", help="Path to input Markdown specification file")
    p_spec_word.add_argument("-o", "--output", help="Output destination Word file path")
    p_spec_word.add_argument("--font", default="TH Sarabun PSK", help="Primary font family")
    p_spec_word.add_argument("-t", "--title", help="Document title")
    p_spec_word.set_defaults(func=cmd_spec_to_word)

    p_spec_docx = spec_subparsers.add_parser("to-docx", help=argparse.SUPPRESS)
    p_spec_docx.add_argument("input_file")
    p_spec_docx.add_argument("-o", "--output")
    p_spec_docx.add_argument("--font", default="TH Sarabun PSK")
    p_spec_docx.add_argument("-t", "--title")
    p_spec_docx.set_defaults(func=cmd_spec_to_word)

    p_spec_audit = spec_subparsers.add_parser("audit", help="Cross-reference project manifest elements with specs")
    p_spec_audit.add_argument("-m", "--manifest", default="project.yaml", help="Path to project manifest")
    p_spec_audit.add_argument("-s", "--specs", help="Path to material specifications manifest or directory")
    strict_group = p_spec_audit.add_mutually_exclusive_group()
    strict_group.add_argument("--strict", dest="strict", action="store_true", help="Exit code 1 if discrepancies exist")
    strict_group.add_argument("--no-strict", dest="strict", action="store_false", default=False)
    p_spec_audit.set_defaults(func=cmd_spec_audit)

    p_spec_add = spec_subparsers.add_parser("add", help="Install specification package into project local specs/")
    p_spec_add.add_argument("pkg_name", help="Package identifier")
    p_spec_add.add_argument("-s", "--specs-dir", default="specs", help="Target project directory")
    p_spec_add.add_argument("-f", "--force", action="store_true", help="Bypass local cache")
    p_spec_add.add_argument("--cache-dir", help="Custom registry cache directory")
    p_spec_add.set_defaults(func=cmd_spec_add)

    p_spec_list = spec_subparsers.add_parser("list", help="List installed specification packages")
    p_spec_list.add_argument("-s", "--specs-dir", default="specs", help="Directory containing installed packages")
    p_spec_list.set_defaults(func=cmd_spec_list)

    # mcp
    p_mcp = subparsers.add_parser("mcp", help="Start Model Context Protocol (MCP) server")
    p_mcp.set_defaults(func=cmd_mcp)

    # docs
    p_docs = subparsers.add_parser("docs", help="Interactive Class Directory documentation engine")
    docs_subparsers = p_docs.add_subparsers(dest="docs_subcommand")

    p_docs_dir = docs_subparsers.add_parser("directory", help="Export interactive BIM Class Directory web app")
    p_docs_dir.add_argument("-o", "--output", default="dist/class_directory.html", help="Output HTML path")
    p_docs_dir.set_defaults(func=cmd_docs_directory)

    p_docs_view = docs_subparsers.add_parser("view", help=argparse.SUPPRESS)
    p_docs_view.add_argument("-o", "--output", default="dist/class_directory.html")
    p_docs_view.set_defaults(func=cmd_docs_directory)

    p_docs_build = docs_subparsers.add_parser("build", help="Generates complete interactive debim Class Directory static site")
    p_docs_build.add_argument("-o", "--output-dir", default="dist/docs", help="Output directory")
    p_docs_build.add_argument("-s", "--schema-version", default="IFC4", help="Target IFC Schema version")
    p_docs_build.set_defaults(func=cmd_docs_build)

    p_docs_serve = docs_subparsers.add_parser("serve", help="Local HTTP server to preview class encyclopedia", description="Spins up a local HTTP server to preview the class encyclopedia in the browser")
    p_docs_serve.add_argument("-d", "--docs-dir", default="dist/docs", help="Directory containing built docs site")
    p_docs_serve.add_argument("-p", "--port", type=int, default=8000, help="Port to serve preview")
    p_docs_serve.add_argument("--no-browser", action="store_true", help="Do not open web browser automatically")
    p_docs_serve.set_defaults(func=cmd_docs_serve)

    p_docs_ex = docs_subparsers.add_parser("exemplar", help="Generate exemplar .debim manifest")
    p_docs_ex.add_argument("-c", "--class-name", help="Target IFC class name or 'ALL'")
    p_docs_ex.add_argument("-o", "--output", help="Path to write generated exemplar YAML")
    p_docs_ex.add_argument("-f", "--format", default="yaml", help="Output format ('yaml' or 'json')")
    p_docs_ex.set_defaults(func=cmd_docs_exemplar)

    p_docs_exp = docs_subparsers.add_parser("export-data", help="Export complete exemplar manifests and Thai spec clauses")
    p_docs_exp.add_argument("-o", "--output-dir", default="dist/docs", help="Target output directory")
    p_docs_exp.set_defaults(func=cmd_docs_export_data)

    return parser


# --- App Wrapper for Typer Compatibility in tests & entrypoints ---

class AppWrapper:
    """Wrapper that emulates typer.Typer interface for backwards-compatibility in existing code/tests."""
    def __call__(self, args: Optional[List[str]] = None):
        return main(args)

    def invoke(self, args: Optional[List[str]] = None):
        return main(args)

app = AppWrapper()


def main(args: Optional[List[str]] = None):
    parser = build_parser()

    if args is None:
        raw_args = sys.argv[1:]
    else:
        raw_args = list(args)

    parsed_args = parser.parse_args(raw_args)

    if parsed_args.version:
        console.print(DEBIM_BANNER)
        console.print(
            f"[bold cyan]debim[/bold cyan] version [bold green]{__version__}[/bold green] "
            f"| Developer: [bold #38bdf8]Prida Takon[/bold #38bdf8] ([dim]@PRIDA-TAKON[/dim])"
        )
        check_and_notify_updates(console=console, quiet=parsed_args.quiet, ctx=parsed_args)
        sys.exit(0)

    if not parsed_args.subcommand:
        console.print(DEBIM_BANNER)
        parser.print_help()
        check_and_notify_updates(console=console, quiet=parsed_args.quiet, ctx=parsed_args)
        sys.exit(0)

    # Check for sub-subcommands if not handled
    if parsed_args.subcommand == "cost" and hasattr(parsed_args, "cost_subcommand") and not parsed_args.cost_subcommand:
        # Ran `debim cost` directly
        pass
    elif parsed_args.subcommand == "scaffold" and not getattr(parsed_args, "scaffold_subcommand", None):
        parser.parse_args([parsed_args.subcommand, "--help"])
        sys.exit(0)
    elif parsed_args.subcommand == "draw" and not getattr(parsed_args, "draw_subcommand", None):
        parser.parse_args([parsed_args.subcommand, "--help"])
        sys.exit(0)
    elif parsed_args.subcommand == "spec" and not getattr(parsed_args, "spec_subcommand", None):
        parser.parse_args([parsed_args.subcommand, "--help"])
        sys.exit(0)
    elif parsed_args.subcommand == "docs" and not getattr(parsed_args, "docs_subcommand", None):
        parser.parse_args([parsed_args.subcommand, "--help"])
        sys.exit(0)

    check_and_notify_updates(console=console, quiet=parsed_args.quiet, ctx=parsed_args)

    if hasattr(parsed_args, "func"):
        parsed_args.func(parsed_args)
    else:
        parser.print_help()
        sys.exit(0)


if __name__ == "__main__":
    main()
