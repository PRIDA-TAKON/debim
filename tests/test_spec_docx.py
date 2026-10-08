"""
Tests for Specification Book MS Word (.docx) export and CLI command.
"""

from pathlib import Path
from typer.testing import CliRunner

from debim.cli import app
from debim.spec.builder import build_specification_book
from debim.spec.docx_exporter import markdown_to_docx, export_spec_book_to_docx

runner = CliRunner()


def test_markdown_to_docx_conversion(tmp_path: Path):
    md_text = """# TEST SPECIFICATION BOOK

**Project Name:** Test Villa
**Date:** 2026-10-08

---

## Table of Contents
1. **Division 03 - Concrete**
   - 1.1 [CONCRETE_240KSC - คอนกรีตผสมเสร็จ 240 ksc](#sec-1)

---

## Detailed Material Specifications

### Division 03: Concrete

#### 1.1 CONCRETE_240KSC - คอนกรีตผสมเสร็จ 240 ksc

- **Specification ID:** `CONCRETE_240KSC`
- **Material Name:** คอนกรีตผสมเสร็จกำลังอัด 240 ksc
- **Manufacturer:** CPAC / SCG

> ข้อควรระวัง: ห้ามเติมน้ำที่หน้างานเด็ดขาด

| ข้อกำหนด | ค่ามาตรฐาน |
|---|---|
| กำลังอัด 28 วัน | 240 ksc |
| ค่ายุบตัว (Slump) | 10 ± 2.5 cm |
"""
    out_docx = tmp_path / "test_spec.docx"
    saved = markdown_to_docx(md_text, out_docx, title="Test Villa Specs")

    assert saved.exists()
    assert saved.stat().st_size > 0

    # Verify python-docx can open and read paragraphs
    import docx
    doc = docx.Document(str(saved))
    texts = [p.text for p in doc.paragraphs if p.text]
    assert any("TEST SPECIFICATION BOOK" in t for t in texts)
    assert any("คอนกรีตผสมเสร็จ" in t for t in texts)
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 3


def test_spec_book_save_docx(tmp_path: Path):
    project_yaml = Path("examples/farnsworth_house/project.yaml")
    specs_yaml = Path("examples/farnsworth_house/specs.yaml")

    book = build_specification_book(
        project_manifest=project_yaml,
        spec_manifest=specs_yaml,
    )

    out_docx = tmp_path / "farnsworth.docx"
    saved_path = book.save(out_docx)
    assert saved_path.exists()
    assert saved_path.stat().st_size > 5000


def test_cli_spec_to_word(tmp_path: Path):
    md_file = tmp_path / "sample.md"
    md_file.write_text("# Spec Title\n\n- Spec clause item\n", encoding="utf-8")
    out_docx = tmp_path / "output.docx"

    res = runner.invoke(app, ["spec", "to-word", str(md_file), "-o", str(out_docx)])
    assert res.exit_code == 0
    assert out_docx.exists()
    assert "MS Word Document Exported Successfully" in res.output
